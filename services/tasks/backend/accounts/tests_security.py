"""Security regression tests (docs/security/tasks-api.md): file links, exports, IDOR and mass-assignment fixes,
public-endpoint throttling, webhook signatures, list-query counts.

Runs on the default SQLite test database; the multi-company counterparts are in tenancy/tests_security.py.
"""
import io
import shutil
import tempfile
from datetime import timedelta
from decimal import Decimal
from unittest import mock

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.models import Role, User
from config import media
from config.exports import safe_cell, safe_row
from crm.models import Lead, Task, TaskAttachment, TaskChangeRequest, TaskStatus


def mk(username, role, department="sales", **extra):
    return User.objects.create_user(username, f"{username}@x.com", "pass@12345", role=role,
                                    department=department, **extra)


class SecurityBase(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._media = tempfile.mkdtemp(prefix="ff-media-")
        cls._override = override_settings(MEDIA_ROOT=cls._media)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._media, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        cache.clear()
        self.admin = mk("boss", Role.ADMIN, "management")
        self.manager = mk("meera", Role.SALES_MANAGER)
        self.rahul = mk("rahul", Role.SALES_EXECUTIVE, reporting_manager=self.manager)
        self.amit = mk("amit", Role.SALES_EXECUTIVE, reporting_manager=self.manager)
        self.rider = mk("ravi", Role.RIDER, "operations")

    def as_(self, user):
        client = APIClient()
        res = client.post("/api/auth/login", {"username": user.username, "password": "pass@12345"})
        client.credentials(HTTP_AUTHORIZATION=f"Bearer {res.data['access']}")
        return client


# --- uploaded files ---------------------------------------------------------------------------------------------

class MediaLinkTests(SecurityBase):
    def setUp(self):
        super().setUp()
        self.task = Task.objects.create(title="Invoice", assigned_to=self.rahul, created_by=self.manager,
                                        due_at=timezone.now() + timedelta(days=1))

    def attach(self, name, body, content_type):
        res = self.as_(self.rahul).post(f"/api/tasks/{self.task.pk}/upload/",
                                        {"file": SimpleUploadedFile(name, body, content_type=content_type)},
                                        format="multipart")
        self.assertEqual(res.status_code, 201, res.content)
        return next(f["url"] for f in res.data if f["filename"] == name)

    def test_uploads_are_not_served_without_a_signed_link(self):
        url = self.attach("invoice.pdf", b"%PDF-1.4 secret", "application/pdf")
        path, _, query = url.partition("?")
        self.assertTrue(query.startswith("s="), url)
        anon = APIClient()
        self.assertEqual(anon.get(path).status_code, 404)                        # guessed path: nothing
        self.assertEqual(anon.get(path + "?s=1:abc").status_code, 404)           # forged signature: nothing
        ok = anon.get(url)                                                       # the link the API gave out
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(b"".join(ok.streaming_content), b"%PDF-1.4 secret")
        self.assertEqual(ok["X-Content-Type-Options"], "nosniff")

    def test_a_signature_is_for_one_file_only(self):
        url = self.attach("a.pdf", b"%PDF-1.4 a", "application/pdf")
        other = self.attach("b.pdf", b"%PDF-1.4 b", "application/pdf")
        token = url.partition("?")[2]
        self.assertEqual(APIClient().get(other.partition("?")[0] + "?" + token).status_code, 404)

    def test_links_expire(self):
        url = self.attach("a.pdf", b"%PDF-1.4 a", "application/pdf")
        with override_settings(MEDIA_LINK_SECONDS=-1):
            self.assertEqual(APIClient().get(url).status_code, 404)

    def test_html_and_svg_uploads_never_render_as_a_page_of_the_app(self):
        for name, ctype in (("x.html", "text/html"), ("x.svg", "image/svg+xml")):
            url = self.attach(name, b"<script>fetch('/api/platform/auth/refresh')</script>", ctype)
            res = APIClient().get(url)
            self.assertEqual(res.status_code, 200)
            self.assertEqual(res["Content-Disposition"], "attachment")
            self.assertIn("sandbox", res["Content-Security-Policy"])

    def test_private_punch_photos_stay_off_media_even_when_signed(self):
        name = "private/attendance/2026/10/x.jpg"
        self.assertEqual(APIClient().get(f"/media/{name}?s={media.sign(name)}").status_code, 404)

    def test_someone_who_cannot_see_the_task_gets_no_link(self):
        self.attach("a.pdf", b"%PDF-1.4 a", "application/pdf")
        self.assertEqual(self.as_(self.rider).get(f"/api/tasks/{self.task.pk}/files/").status_code, 404)
        self.assertTrue(TaskAttachment.objects.filter(task=self.task).exists())

    def test_leave_documents_have_a_size_limit(self):
        from hr.models import LeaveType
        lt = LeaveType.objects.create(name="Sick leave")
        big = SimpleUploadedFile("cert.pdf", b"0" * (10 * 1024 * 1024 + 1), content_type="application/pdf")
        start = timezone.localdate() + timedelta(days=7)
        while start.weekday() >= 5:
            start += timedelta(days=1)
        res = self.as_(self.rahul).post("/api/leaves/", {"leave_type": lt.pk, "start_date": start, "end_date": start,
                                                         "reason": "fever", "document": big}, format="multipart")
        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn("document", res.data)


class ImageBombTests(TestCase):
    def test_a_huge_png_canvas_is_never_decoded(self):
        from PIL import Image

        from config.images import compress_image
        buf = io.BytesIO()
        Image.new("1", (8000, 8000)).save(buf, "PNG")  # 64 MP from a few KB
        raw = buf.getvalue()
        self.assertLess(len(raw), 200_000)
        out, ext = compress_image(raw)
        self.assertIsNone(ext)
        self.assertIs(out, raw)


# --- exports ----------------------------------------------------------------------------------------------------

class FormulaInjectionTests(SecurityBase):
    def test_cells_that_start_a_formula_are_written_as_text(self):
        self.assertEqual(safe_cell("=HYPERLINK(\"http://x\")"), "'=HYPERLINK(\"http://x\")")
        for bad in ("+1", "-2+3", "@SUM(A1)", "\tx", "\rx"):
            self.assertTrue(safe_cell(bad).startswith("'"), bad)
        self.assertEqual(safe_row([5, None, "plain", Decimal("1.5")]), [5, None, "plain", Decimal("1.5")])

    def test_attendance_reports_neutralise_typed_text(self):
        from openpyxl import load_workbook

        from hr.team import to_csv, to_xlsx
        rows = [["Rahul", "=cmd|' /C calc'!A0"]]
        self.assertIn(b"'=cmd", to_csv(["Name", "Reason"], rows))
        ws = load_workbook(io.BytesIO(to_xlsx("Leave", ["Name", "Reason"], rows))).active
        self.assertEqual(ws["B2"].value, "'=cmd|' /C calc'!A0")
        self.assertNotEqual(ws["B2"].data_type, "f")


@override_settings(DEBUG=False, TESTING=False)
class WhatsAppWebhookTests(TestCase):
    def test_unsigned_posts_are_refused_in_production(self):
        with mock.patch.dict("os.environ", {"WHATSAPP_APP_SECRET": ""}):
            res = APIClient().post("/api/webhooks/whatsapp", {"entry": []}, format="json")
        self.assertEqual(res.status_code, 403)

    def test_signed_posts_are_accepted(self):
        import hashlib
        import hmac
        import json
        body = json.dumps({"entry": []}).encode()
        sig = "sha256=" + hmac.new(b"s3cret", body, hashlib.sha256).hexdigest()
        with mock.patch.dict("os.environ", {"WHATSAPP_APP_SECRET": "s3cret"}):
            ok = APIClient().generic("POST", "/api/webhooks/whatsapp", body, content_type="application/json",
                                     HTTP_X_HUB_SIGNATURE_256=sig)
            bad = APIClient().generic("POST", "/api/webhooks/whatsapp", body, content_type="application/json",
                                      HTTP_X_HUB_SIGNATURE_256="sha256=ünicode")
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(bad.status_code, 403)


class LegacyLoginOffInMultiCompanyModeTests(SecurityBase):
    def test_a_legacy_token_is_ignored_when_companies_are_on(self):
        client = self.as_(self.admin)
        self.assertEqual(client.get("/api/auth/me").status_code, 200)
        from accounts.platform_auth import LegacyJWTAuthentication
        request = mock.Mock(META={"HTTP_AUTHORIZATION": client._credentials["HTTP_AUTHORIZATION"]})
        with mock.patch("tenancy.schemas.enabled", return_value=True):
            self.assertIsNone(LegacyJWTAuthentication().authenticate(request))


# --- tasks and leads --------------------------------------------------------------------------------------------

class TaskAbuseTests(SecurityBase):
    def due(self):
        return (timezone.now() + timedelta(days=2)).isoformat()

    def test_a_task_cannot_be_created_already_done(self):
        res = self.as_(self.manager).post("/api/tasks/", {"title": "Pad my score", "assigned_to": self.rahul.pk,
                                                       "status": "done", "effort_minutes": 600,
                                                       "due_at": self.due()}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(Task.objects.get(pk=res.data["id"]).status, TaskStatus.OPEN)

    def test_restore_only_reaches_tasks_in_the_callers_scope(self):
        from accounts.access import forget
        from accounts.models import AccessGrant, AccessRole
        hidden = Task.objects.create(title="Purchase secret", assigned_to=self.admin, created_by=self.admin,
                                     deleted_at=timezone.now())
        role = AccessRole.objects.get(code=Role.SALES_MANAGER)
        AccessGrant.objects.get_or_create(role=role, permission="tasks.bin", defaults={"scope": ""})
        forget()
        res = self.as_(self.manager).post(f"/api/tasks/{hidden.pk}/restore/")
        self.assertEqual(res.status_code, 400)
        hidden.refresh_from_db()
        self.assertIsNotNone(hidden.deleted_at)

    def test_change_request_with_a_naive_deadline_keeps_the_inbox_working(self):
        task = Task.objects.create(title="Quote", assigned_to=self.rahul, created_by=self.manager,
                                   due_at=timezone.now() + timedelta(days=1))
        naive = (timezone.localtime() + timedelta(days=5)).replace(tzinfo=None).strftime("%Y-%m-%dT%H:%M")
        res = self.as_(self.rahul).post(f"/api/tasks/{task.pk}/request_change/",
                                        {"changes": {"due_at": naive}, "reason": "customer moved it"}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        self.assertIn("+", TaskChangeRequest.objects.get().changes["due_at"])
        self.assertEqual(self.as_(self.manager).get("/api/task-change-requests/?scope=inbox").status_code, 200)

    def test_change_request_values_are_checked_before_saving(self):
        task = Task.objects.create(title="Quote", assigned_to=self.rahul, created_by=self.manager,
                                   due_at=timezone.now() + timedelta(days=1))
        for changes in ({"assigned_to": "x"}, {"priority": "<script>"}, {"due_at": "soon"},
                        {"repeat_until": "never"}):
            res = self.as_(self.rahul).post(f"/api/tasks/{task.pk}/request_change/",
                                            {"changes": changes, "reason": "because"}, format="json")
            self.assertEqual(res.status_code, 400, (changes, res.content))
        self.assertFalse(TaskChangeRequest.objects.exists())

    def test_paid_ai_calls_are_rate_limited_per_person(self):
        client = self.as_(self.rahul)
        with mock.patch("rest_framework.throttling.SimpleRateThrottle.THROTTLE_RATES", {"ai": "1/min"}):
            first = client.post("/api/tasks/proofread/", {"text": "helo"}, format="json").status_code
            second = client.post("/api/tasks/proofread/", {"text": "helo"}, format="json").status_code
            other = self.as_(self.amit).post("/api/tasks/proofread/", {"text": "helo"}, format="json").status_code
        self.assertNotEqual(first, 429)
        self.assertEqual(second, 429)
        self.assertNotEqual(other, 429)

    def test_a_role_without_leads_cannot_create_leads(self):
        res = self.as_(self.rider).post("/api/leads/", {"customer_name": "Spam", "department": "sales"}, format="json")
        self.assertEqual(res.status_code, 403)
        self.assertFalse(Lead.objects.filter(customer_name="Spam").exists())

    def test_task_list_queries_do_not_grow_with_the_page(self):
        client = self.as_(self.admin)

        def count(n):
            Task.objects.all().delete()
            for i in range(n):
                t = Task.objects.create(title=f"T{i}", assigned_to=(self.rahul, self.amit, self.manager)[i % 3],
                                        created_by=self.admin, due_at=timezone.now() + timedelta(days=1))
                t.subscribers.add(self.admin)
            with CaptureQueriesContext(connection) as ctx:
                res = client.get("/api/tasks/?page_size=100")
            self.assertEqual(res.status_code, 200)
            self.assertEqual(res.data["count"], n)
            return len(ctx.captured_queries)

        few, many = count(3), count(30)
        self.assertLessEqual(many - few, 3, f"{few} queries for 3 tasks, {many} for 30")


# --- mistakes ---------------------------------------------------------------------------------------------------

class MistakeTests(SecurityBase):
    def setUp(self):
        super().setUp()
        from mistakes.models import Mistake, MistakeCategory
        MistakeCategory.objects.get_or_create(name="Billing", defaults={"active": True})
        self.own = Mistake.objects.create(employee=self.rahul, reported_by=self.rahul, category="Billing",
                                          severity="high", description="Wrong invoice", department="sales",
                                          manager=self.manager)

    def test_the_employee_cannot_rewrite_or_move_a_logged_mistake(self):
        client = self.as_(self.rahul)
        self.assertEqual(client.patch(f"/api/mistakes/{self.own.pk}/", {"severity": "low"}, format="json").status_code, 403)
        self.assertEqual(client.patch(f"/api/mistakes/{self.own.pk}/", {"employee": self.amit.pk}, format="json").status_code, 403)
        self.own.refresh_from_db()
        self.assertEqual((self.own.employee_id, self.own.severity), (self.rahul.pk, "high"))

    def test_a_reviewer_still_cannot_move_it_to_someone_else(self):
        res = self.as_(self.manager).patch(f"/api/mistakes/{self.own.pk}/", {"employee": self.amit.pk}, format="json")
        self.assertEqual(res.status_code, 400)
        ok = self.as_(self.manager).patch(f"/api/mistakes/{self.own.pk}/", {"severity": "medium"}, format="json")
        self.assertEqual(ok.status_code, 200, ok.content)

    def test_suggestions_naming_other_people_are_for_reviewers_only(self):
        self.assertEqual(self.as_(self.rahul).get(f"/api/mistakes/{self.own.pk}/suggestions/").status_code, 403)

    def test_nobody_reviews_their_own_mistake_whatever_their_reach(self):
        from mistakes.models import Mistake
        from mistakes.views import can_review
        mine = Mistake.objects.create(employee=self.admin, reported_by=self.manager, category="Billing",
                                      description="x", department="management")
        self.assertFalse(can_review(self.admin, mine))

    def test_a_linked_task_must_be_visible(self):
        hidden = Task.objects.create(title="Purchase secret", assigned_to=self.admin, created_by=self.admin)
        res = self.as_(self.rahul).post("/api/mistakes/", {"employee": self.rahul.pk, "category": "Billing",
                                                           "description": "x", "task": hidden.pk}, format="json")
        self.assertEqual(res.status_code, 400, res.content)


# --- HR and payroll ---------------------------------------------------------------------------------------------

class LeaveAndCorrectionTests(SecurityBase):
    def setUp(self):
        super().setUp()
        from hr.models import LeaveType
        # A quota big enough for the 30-day stretch below (the balance is enforced since PLAN.md §19.33).
        self.lt = LeaveType.objects.create(name="Casual leave", annual_quota=60)
        day = timezone.localdate() + timedelta(days=14)
        while day.weekday() >= 5:
            day += timedelta(days=1)
        self.day = day

    def apply(self, user):
        res = self.as_(user).post("/api/leaves/", {"leave_type": self.lt.pk, "start_date": self.day,
                                                   "end_date": self.day, "reason": "family"}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        return res.data["id"]

    def test_stretching_a_pending_leave_recounts_its_days(self):
        from hr.models import LeaveRequest
        pk = self.apply(self.rahul)
        res = self.as_(self.rahul).patch(f"/api/leaves/{pk}/", {"end_date": self.day + timedelta(days=30)}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        self.assertGreater(LeaveRequest.objects.get(pk=pk).days, 15)

    def test_an_approver_cannot_edit_someone_elses_leave(self):
        pk = self.apply(self.rahul)
        res = self.as_(self.manager).patch(f"/api/leaves/{pk}/?scope=team", {"end_date": self.day + timedelta(days=9)},
                                           format="json")
        self.assertEqual(res.status_code, 403)

    def test_a_reviewed_correction_cannot_be_rewritten(self):
        from hr.models import AttendanceCorrection
        when = timezone.now() - timedelta(days=1)
        c = AttendanceCorrection.objects.create(user=self.rahul, date=when.date(), requested_check_in=when,
                                                reason="forgot", status="approved")
        res = self.as_(self.rahul).patch(f"/api/attendance-corrections/{c.pk}/",
                                         {"requested_check_in": (when - timedelta(hours=3)).isoformat()}, format="json")
        self.assertEqual(res.status_code, 403)

    def test_a_correction_cannot_be_for_a_future_day(self):
        later = timezone.now() + timedelta(days=3)
        res = self.as_(self.rahul).post("/api/attendance-corrections/", {
            "date": later.date().isoformat(), "requested_check_in": later.isoformat(), "reason": "x"}, format="json")
        self.assertEqual(res.status_code, 400, res.content)

    def test_attendance_import_reads_a_bounded_number_of_rows(self):
        from rest_framework.exceptions import ValidationError

        from hr import team
        body = "Employee ID,Date,Status\n" + "\n".join(f"E{i},2026-10-01,P" for i in range(team.IMPORT_MAX_ROWS * 3))
        upload = SimpleUploadedFile("big.csv", body.encode(), content_type="text/csv")
        with mock.patch.object(team, "IMPORT_MAX_BYTES", 10 ** 9):
            with self.assertRaises(ValidationError):
                team.read_rows(upload)


class PayrollTests(SecurityBase):
    def test_a_finalised_run_cannot_be_edited(self):
        from payroll.models import PayrollRun
        run = PayrollRun.objects.create(year=2026, month=9, status="finalised")
        res = self.as_(self.admin).patch(f"/api/payroll-runs/{run.pk}/", {"month": 8}, format="json")
        self.assertEqual(res.status_code, 400, res.content)
        run.refresh_from_db()
        self.assertEqual(run.month, 9)


# --- users ------------------------------------------------------------------------------------------------------

class UserAdminTests(SecurityBase):
    def test_hr_cannot_switch_an_admin_back_on(self):
        hr = mk("hema", Role.HR_MANAGER, "hr")
        other_admin = mk("adi", Role.ADMIN, "management", is_active=False)
        res = self.as_(hr).post(f"/api/users/{other_admin.pk}/activate/")
        self.assertEqual(res.status_code, 403)
        other_admin.refresh_from_db()
        self.assertFalse(other_admin.is_active)

    def test_hr_cannot_edit_someone_whose_platform_role_is_admin(self):
        from accounts.access import forget
        from accounts.models import AccessRole, UserAccessRole
        hr = mk("hema", Role.HR_MANAGER, "hr")
        UserAccessRole.objects.create(user=self.rahul, role=AccessRole.objects.get(code=Role.ADMIN))
        forget()
        res = self.as_(hr).patch(f"/api/users/{self.rahul.pk}/", {"password": "Takeover@123"}, format="json")
        self.assertEqual(res.status_code, 403)
