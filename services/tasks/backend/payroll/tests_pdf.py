"""Payslips as PDF and payroll downloads (PLAN.md §19.31)."""
import io
import zipfile
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from accounts.models import Role

from . import services
from .models import Advance, SalaryStructure
from .pdf import amount_in_words, indian, inr, payslip_pdf
from .tests import MONTH, YEAR, Base, make, mark


def _text(pdf: bytes) -> str:
    """The page text of a reportlab PDF (its content streams are zlib-compressed)."""
    import re
    import zlib
    out = []
    for m in re.finditer(rb"stream\r?\n(.*?)endstream", pdf, re.S):
        raw = m.group(1).strip()
        try:   # /Filter [ /ASCII85Decode /FlateDecode ]
            import base64
            raw = zlib.decompress(base64.a85decode(raw[:-2] if raw.endswith(b"~>") else raw))
        except Exception:   # noqa: BLE001 -- images, fonts
            continue
        out.append(raw.decode("latin-1"))
    return "\n".join(out)


class IndianFormatTests(Base):
    def test_lakh_grouping(self):
        self.assertEqual(indian(0), "0.00")
        self.assertEqual(indian(999), "999.00")
        self.assertEqual(indian(1000), "1,000.00")
        self.assertEqual(indian(123456.5), "1,23,456.50")
        self.assertEqual(indian(Decimal("12345678.9")), "1,23,45,678.90")
        self.assertEqual(inr(26000), "₹26,000.00")

    def test_amount_in_words(self):
        self.assertEqual(amount_in_words(0), "Rupees Zero Only")
        self.assertEqual(amount_in_words(26000), "Rupees Twenty-Six Thousand Only")
        self.assertEqual(amount_in_words(Decimal("123456.50")),
                         "Rupees One Lakh Twenty-Three Thousand Four Hundred Fifty-Six and Fifty Paise Only")
        self.assertEqual(amount_in_words(12345678), "Rupees One Crore Twenty-Three Lakh Forty-Five Thousand "
                                                    "Six Hundred Seventy-Eight Only")


class PayslipPdfTests(Base):
    def setUp(self):
        super().setUp()
        for day in range(1, 31):
            if date(YEAR, MONTH, day).weekday() != 6:
                mark(self.rahul, day)
        self.rahul.first_name, self.rahul.last_name = "Rahul", "Verma"
        self.rahul.employee_code, self.rahul.designation = "CT-0042", "Sales executive"
        self.rahul.save()
        Advance.objects.create(user=self.rahul, amount=Decimal("1500"), given_on=date(2026, 8, 20))
        self.other = make("arjun", Role.SALES_EXECUTIVE)
        SalaryStructure.objects.create(user=self.other, monthly_gross=Decimal("30000"), effective_from=date(2026, 1, 1))
        self.run = self.run_payroll()
        self.slip = self.run.payslips.get(user=self.rahul)
        self.other_slip = self.run.payslips.get(user=self.other)

    def test_content(self):
        pdf = payslip_pdf([self.slip])
        self.assertTrue(pdf.startswith(b"%PDF"))
        text = _text(pdf)
        for want in ("Rahul Verma", "CT-0042", "Sales executive", "Payslip for September 2026",
                     "26,000.00", "Advance recovery", "1,500.00", "24,500.00",
                     "Rupees Twenty-Four Thousand Five Hundred Only", "DRAFT"):
            self.assertIn(want, text, want)

    def test_own_payslip_only_when_finalised(self):
        self.as_(self.rahul)
        url = f"/api/payslips/{self.slip.id}/pdf/"
        self.assertEqual(self.client.get(url).status_code, 404)   # draft: not visible, so no PDF either
        services.finalise(self.run)
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/pdf")
        self.assertIn("payslip-2026-09-rahul-verma.pdf", res["Content-Disposition"])
        self.assertTrue(res.content.startswith(b"%PDF"))
        self.assertNotIn("DRAFT", _text(res.content))

    def test_others_payslip_forbidden_for_a_member(self):
        services.finalise(self.run)
        self.as_(self.rahul)
        self.assertEqual(self.client.get(f"/api/payslips/{self.other_slip.id}/pdf/").status_code, 404)
        self.as_(self.manager)   # a sales manager has no payroll rights at all
        self.assertEqual(self.client.get(f"/api/payslips/{self.slip.id}/pdf/").status_code, 404)

    def test_payroll_manager_downloads_anyones_and_the_whole_run(self):
        self.as_(self.hr)
        res = self.client.get(f"/api/payslips/{self.other_slip.id}/pdf/")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.content.startswith(b"%PDF"))
        z = self.client.get(f"/api/payroll-runs/{self.run.id}/payslips_pdf/")
        self.assertEqual(z.status_code, 200)
        self.assertEqual(z["Content-Type"], "application/zip")
        names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
        self.assertEqual(len(names), self.run.payslips.count())
        self.assertIn("payslip-2026-09-rahul-verma.pdf", names)
        one = self.client.get(f"/api/payroll-runs/{self.run.id}/payslips_pdf/?bundle=pdf")
        self.assertEqual(one["Content-Type"], "application/pdf")
        self.assertEqual(one.content.count(b"/Type /Page\n") + one.content.count(b"/Type /Page "),
                         self.run.payslips.count())

    def test_whole_run_needs_payroll_run(self):
        self.as_(self.rahul)
        self.assertEqual(self.client.get(f"/api/payroll-runs/{self.run.id}/payslips_pdf/").status_code, 403)


class PayrollSheetTests(Base):
    def test_run_export_xlsx_and_lists(self):
        for day in range(1, 31):
            if date(YEAR, MONTH, day).weekday() != 6:
                mark(self.rahul, day)
        Advance.objects.create(user=self.rahul, amount=Decimal("500"), given_on=date(2026, 8, 1), reason="=HYPERLINK(\"x\")")
        run = self.run_payroll()
        self.as_(self.hr)
        res = self.client.get(f"/api/payroll-runs/{run.id}/export/?export=xlsx")
        self.assertEqual(res.status_code, 200)
        ws = load_workbook(io.BytesIO(res.content)).active
        self.assertEqual(ws["A1"].value, "Employee")
        self.assertIn("rahul", [r[1] for r in ws.iter_rows(min_row=2, values_only=True)])
        # runs, salaries, advances, payslips lists
        self.assertEqual(self.client.get("/api/payroll-runs/?export=csv").status_code, 200)
        sal = self.client.get(f"/api/salary-structures/?user={self.rahul.id}&export=csv").content.decode("utf-8-sig")
        self.assertEqual(len(sal.strip().splitlines()), 2)
        adv = self.client.get("/api/advances/?export=csv").content.decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", adv)   # formula neutralised
        slips = load_workbook(io.BytesIO(self.client.get(f"/api/payslips/?user={self.rahul.id}&export=xlsx").content)).active
        self.assertEqual(slips.max_row, 2)

    def test_member_exports_only_own(self):
        services.finalise(self.run_payroll())
        self.as_(self.rahul)
        rows = self.client.get("/api/payslips/?export=csv").content.decode("utf-8-sig").strip().splitlines()
        self.assertEqual(len(rows), 2)
        self.assertEqual(self.client.get("/api/payroll-runs/?export=csv").status_code, 403)
