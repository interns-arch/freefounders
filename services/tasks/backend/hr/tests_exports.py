"""Attendance downloads (PLAN.md §19.31): ?export=xlsx|csv on the attendance lists has every matching row, the
list's own scope, and never a live formula."""
import csv
import io
from datetime import date

from openpyxl import load_workbook

from hr.models import Attendance, AttendanceCorrection, LeaveRequest, LeaveStatus, LeaveType
from hr.tests_team_views import MON, SEP, SUN, _Base, at


def sheet(res):
    assert res.status_code == 200, res.content
    assert res["Content-Type"].startswith("application/vnd.openxmlformats"), res["Content-Type"]
    ws = load_workbook(io.BytesIO(res.content)).active
    return [[c.value for c in r] for r in ws.iter_rows()]


def csv_rows(res):
    assert res.status_code == 200, res.content
    return list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))


class ExportTests(_Base):
    def setUp(self):
        super().setUp()
        Attendance.objects.create(user=self.rahul, date=MON, status="late", check_in=at(MON, 10, 5),
                                  check_out=at(MON, 18, 0), working_minutes=475, is_late=True,
                                  note="=HYPERLINK(\"http://x\")")
        Attendance.objects.create(user=self.priya, date=MON, status="present", check_in=at(MON, 9, 0),
                                  check_out=at(MON, 18, 0), working_minutes=540)
        self.cl = LeaveType.objects.create(name="Casual", annual_quota=12)

    # ---- my month
    def test_my_month_one_row_per_day(self):
        rows = sheet(self.as_(self.rahul).get("/api/attendance/monthly/", {"year": SEP[0], "month": SEP[1], "export": "xlsx"}))
        self.assertEqual(rows[0][:4], ["Date", "Day", "Employee", "Status"])
        self.assertEqual(len(rows), 1 + 30)
        mon = next(r for r in rows[1:] if r[0].date() == MON)
        self.assertEqual(mon[3], "Late")
        self.assertEqual(next(r for r in rows[1:] if r[0].date() == SUN)[3], "Week off")

    def test_member_cannot_export_someone_elses_month(self):
        res = self.as_(self.rahul).get("/api/attendance/monthly/", {"year": SEP[0], "month": SEP[1], "user": self.priya.id, "export": "csv"})
        self.assertEqual(res.status_code, 403)

    # ---- live / daily
    def test_team_day_filters_and_scope(self):
        hr = self.as_(self.hr)
        rows = csv_rows(hr.get("/api/attendance/team_day/", {"date": MON.isoformat(), "export": "csv"}))
        names = {r[2] for r in rows[1:]}
        self.assertTrue({"Rahul Verma", "Priya"} <= names)
        # the status tile (late) and department narrow the file like the screen
        late = csv_rows(hr.get("/api/attendance/team_day/", {"date": MON.isoformat(), "export": "csv", "filter": "late"}))
        self.assertEqual([r[2] for r in late[1:]], ["Rahul Verma"])
        dep = csv_rows(hr.get("/api/attendance/team_day/", {"date": MON.isoformat(), "export": "csv", "department": "purchase"}))
        self.assertEqual([r[2] for r in dep[1:]], ["Priya"])
        # a sales manager sees only their department
        mgr = csv_rows(self.as_(self.manager).get("/api/attendance/team_day/", {"date": MON.isoformat(), "export": "csv"}))
        self.assertNotIn("Priya", {r[2] for r in mgr[1:]})
        self.assertEqual(self.as_(self.rahul).get("/api/attendance/team_day/", {"export": "csv"}).status_code, 403)

    def test_formula_is_neutralised(self):
        rows = sheet(self.as_(self.hr).get("/api/attendance/team_day/", {"date": MON.isoformat(), "export": "xlsx", "filter": "late"}))
        self.assertEqual(rows[1][-1], "'=HYPERLINK(\"http://x\")")

    # ---- dashboard
    def test_team_month_sorted_and_filtered(self):
        hr = self.as_(self.hr)
        rows = sheet(hr.get("/api/attendance/team_month/", {"year": SEP[0], "month": SEP[1], "export": "xlsx",
                                                          "sort": "hours", "dir": "desc"}))
        self.assertEqual(rows[0][:2], ["Employee ID", "Name"])
        self.assertEqual(rows[1][1], "Priya")       # 540 minutes first
        self.assertEqual(rows[2][1], "Rahul Verma")
        q = sheet(hr.get("/api/attendance/team_month/", {"year": SEP[0], "month": SEP[1], "export": "xlsx", "q": "rahul"}))
        self.assertEqual([r[1] for r in q[1:]], ["Rahul Verma"])
        self.assertEqual(self.as_(self.rahul).get("/api/attendance/team_month/", {"export": "xlsx"}).status_code, 403)

    # ---- leave + corrections
    def test_leave_exports_all_rows_and_scope(self):
        for i in range(30):   # more than one page (25)
            LeaveRequest.objects.create(user=self.rahul, leave_type=self.cl, start_date=date(2026, 10, 1),
                                        end_date=date(2026, 10, 1), days=1, reason=f"r{i}")
        LeaveRequest.objects.create(user=self.priya, leave_type=self.cl, start_date=date(2026, 10, 2),
                                    end_date=date(2026, 10, 2), days=1, reason="+cmd", status=LeaveStatus.APPROVED)
        mine = sheet(self.as_(self.rahul).get("/api/leaves/", {"export": "xlsx"}))
        self.assertEqual(len(mine), 31)
        self.assertEqual(mine[0][:3], ["Employee", "Username", "Leave type"])
        self.assertTrue(all(r[1] == "rahul" for r in mine[1:]))
        team = csv_rows(self.as_(self.hr).get("/api/leaves/", {"export": "csv", "scope": "team", "status": "approved"}))
        self.assertEqual(len(team), 2)
        self.assertEqual(team[1][6], "'+cmd")

    def test_corrections_pending_for_approver_only(self):
        AttendanceCorrection.objects.create(user=self.rahul, date=MON, reason="forgot")
        AttendanceCorrection.objects.create(user=self.rahul, date=SUN, reason="old", status="approved")
        pending = csv_rows(self.as_(self.hr).get("/api/attendance-corrections/", {"export": "csv", "scope": "team", "status": "pending"}))
        self.assertEqual(len(pending), 2)
        self.assertEqual(pending[1][0], "Rahul Verma")
        # a member asking for "team" only gets their own rows
        own = csv_rows(self.as_(self.priya).get("/api/attendance-corrections/", {"export": "csv", "scope": "team"}))
        self.assertEqual(len(own), 1)
