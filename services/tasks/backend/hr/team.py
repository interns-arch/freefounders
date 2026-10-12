"""Attendance for managers (PLAN.md §19.15): everyone's day, everyone's month, report files and the
attendance import. Inspired by SalaryBox's Live / Daily attendance, Attendance Dashboard and Reports.

Scope (PLAN.md §19.19): each screen is its own permission with a scope -- Live attendance (hr.live), Daily
attendance (hr.daily), Attendance dashboard (hr.dashboard), Attendance reports (hr.reports): Team = me + my direct
reports, Department = my department + my team, Company = everyone. Under the "manager approves" rule a reports-to
manager also sees their own team on these screens. Anyone else -> 403. Everything runs inside the request's company
schema (TENANCY=schemas) like the rest of HR.

The day classification is the one `services.monthly_report` uses: a stored row wins; otherwise holiday,
week-off, "not checked in" for today, absent for a past working day, nothing for a future day.
"""
import csv
import io
import re
from itertools import islice
from datetime import date, datetime, time, timedelta

from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from rest_framework.exceptions import PermissionDenied, ValidationError

from accounts.models import Department, DepartmentOption, User
from accounts.permissions import has_capability
from crm.models import Holiday

from . import approvals, services
from .models import Attendance, AttendanceStatus, LeaveRequest

STATUSES = [s.value for s in AttendanceStatus]
MAX_RANGE_DAYS = 366
IMPORT_MAX_BYTES = 2 * 1024 * 1024
IMPORT_MAX_ROWS = 5000
IMPORT_MAX_COLUMNS = 64


# ------------------------------------------------------------------ scope

# The managers' attendance screens; each is its own permission with a scope.
VIEW_KEYS = ("hr.live", "hr.daily", "hr.dashboard", "hr.reports")


def screen_scope(user, keys=VIEW_KEYS) -> str | None:
    """How far this user sees on these screens: the widest scope of the given permissions; under the "manager
    approves" rule a reports-to manager sees at least their team. None: not at all."""
    from accounts.access import scope_of, widest
    scope = widest(*(scope_of(user, k) for k in keys))
    if approvals.mode() == "manager" and user.reports.filter(is_active=True).exists():
        scope = widest(scope, "team")
    return scope


def people_in_scope(user, scope, base=None):
    """Team = me + my direct reports; Department = my department + my team; Company = everyone; None = only me."""
    base = User.objects.filter(is_active=True) if base is None else base
    if scope == "company":
        return base
    cond = Q(pk=user.pk)
    if scope in ("team", "department"):
        cond |= Q(reporting_manager=user)
    if scope == "department":
        cond |= Q(department=user.department)
    return base.filter(cond).distinct()


def sees_company(user, keys=VIEW_KEYS) -> bool:
    return screen_scope(user, keys) == "company"


def may_view_team(user, keys=VIEW_KEYS) -> bool:
    return screen_scope(user, keys) is not None


def team_users(user, *, inactive=False, department="", q="", keys=VIEW_KEYS):
    """The people this manager may see on these screens, filtered. Raises PermissionDenied without the permission."""
    base = User.objects.all() if inactive else User.objects.filter(is_active=True)
    scope = screen_scope(user, keys)
    if scope is None:
        raise PermissionDenied("Your role does not include this attendance screen.")
    qs = people_in_scope(user, scope, base)
    if department:
        qs = qs.filter(department=department)
    q = (q or "").strip()
    if q:
        cond = Q(username__icontains=q) | Q(employee_code__icontains=q) | Q(whatsapp_phone__icontains=q)
        for word in q.split():
            cond |= Q(first_name__icontains=word) | Q(last_name__icontains=word)
        qs = qs.filter(cond)
    return list(qs.distinct().order_by("first_name", "last_name", "username"))


def department_names() -> dict:
    names = dict(Department.choices)
    names.update({d.code: d.name for d in DepartmentOption.objects.all()})
    return names


def departments_of(users, names) -> list:
    codes = sorted({u.department for u in users if u.department}, key=lambda c: names.get(c, c).lower())
    return [{"code": c, "name": names.get(c, c)} for c in codes]


def person(u, names) -> dict:
    return {
        "user_id": u.id,
        "name": u.get_full_name() or u.username,
        "username": u.username,
        "employee_code": u.employee_code,
        "phone": u.whatsapp_phone,
        "department": names.get(u.department, u.department or ""),
        "department_code": u.department,
        "designation": u.designation,
        "is_active": u.is_active,
    }


# ---------------------------------------------------------- classification

class Calendar:
    """Holidays and week-offs for a range, read once (no query per day)."""

    def __init__(self, first: date, last: date):
        self.holidays = {h.date: h.name for h in Holiday.objects.filter(date__range=(first, last))}
        self.week_offs = set(services.cfg()["week_offs"])
        self.today = timezone.localdate()

    def status_without_record(self, day: date):
        if day in self.holidays:
            return AttendanceStatus.HOLIDAY.value
        if day.weekday() in self.week_offs:
            return AttendanceStatus.WEEK_OFF.value
        if day > self.today:
            return None
        if day == self.today:
            return "not_checked_in"
        return AttendanceStatus.ABSENT.value


def parse_day(text, field="date") -> date:
    try:
        return datetime.strptime(str(text), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        raise ValidationError({field: "Use YYYY-MM-DD."})


def parse_month(params):
    today = timezone.localdate()
    try:
        year = int(params.get("year") or today.year)
        month = int(params.get("month") or today.month)
        if not (1 <= month <= 12 and 2000 <= year <= 2100):
            raise ValueError
    except ValueError:
        raise ValidationError({"month": "Invalid year/month."})
    first = date(year, month, 1)
    last = date(year + (month == 12), (month % 12) + 1, 1) - timedelta(days=1)
    return year, month, first, last


def parse_range(params):
    first = parse_day(params.get("from"), "from")
    last = parse_day(params.get("to"), "to")
    if last < first:
        raise ValidationError({"to": "The end date is before the start date."})
    if (last - first).days >= MAX_RANGE_DAYS:
        raise ValidationError({"to": f"Pick at most {MAX_RANGE_DAYS} days."})
    return first, last


def team_day(users, day: date) -> dict:
    """One row per person with their status for the day."""
    names = department_names()
    cal = Calendar(day, day)
    records = {a.user_id: a for a in Attendance.objects.filter(user__in=[u.id for u in users], date=day)}
    rows = []
    for u in users:
        rec = records.get(u.id)
        row = person(u, names)
        if rec:
            row.update({
                "status": rec.status,
                "check_in": rec.check_in, "check_out": rec.check_out,
                "working_minutes": rec.working_minutes,
                "is_late": rec.is_late, "is_early_checkout": rec.is_early_checkout,
                "missing_checkout": rec.missing_checkout,
                "attendance_id": rec.id, "note": rec.note,
                "has_check_in_photo": bool(rec.check_in_photo),
                "has_check_out_photo": bool(rec.check_out_photo),
            })
        else:
            row.update({
                "status": cal.status_without_record(day),
                "check_in": None, "check_out": None, "working_minutes": 0,
                "is_late": False, "is_early_checkout": False, "missing_checkout": False,
                "attendance_id": None, "note": "",
                "has_check_in_photo": False, "has_check_out_photo": False,
            })
        rows.append(row)
    return {
        "date": day.isoformat(),
        "is_today": day == cal.today,
        "holiday": cal.holidays.get(day),
        "week_off": day.weekday() in cal.week_offs,
        "rows": rows,
        "departments": departments_of(users, names),
    }


def month_statuses(users, first: date, last: date):
    """{user_id: {date: (status, record|None)}} for the range, with three queries in all."""
    cal = Calendar(first, last)
    by_user = {}
    for a in Attendance.objects.filter(user__in=[u.id for u in users], date__range=(first, last)):
        by_user.setdefault(a.user_id, {})[a.date] = a
    out = {}
    for u in users:
        recs = by_user.get(u.id, {})
        days, cur = {}, first
        while cur <= last:
            rec = recs.get(cur)
            days[cur] = (rec.status, rec) if rec else (cal.status_without_record(cur), None)
            cur += timedelta(days=1)
        out[u.id] = days
    return out, cal


def totals_for(days: dict) -> dict:
    """Totals in the shape of monthly_report's `totals`, plus late marks and hours."""
    totals = {s: 0 for s in STATUSES}
    minutes = late_marks = 0
    for status, rec in days.values():
        if status in totals:          # "not_checked_in" (today) and future days are not counted
            totals[status] += 1
        if rec:
            minutes += rec.working_minutes
            late_marks += 1 if rec.is_late else 0
    return {
        "totals": totals,
        "present_days": totals["present"] + totals["late"],
        "late_marks": late_marks,
        "working_minutes": minutes,
        "total_hours": round(minutes / 60, 1),
    }


def team_month(users, year, month, first, last) -> dict:
    names = department_names()
    statuses, _ = month_statuses(users, first, last)
    rows = []
    for u in users:
        row = person(u, names)
        row.update(totals_for(statuses[u.id]))
        rows.append(row)
    return {"year": year, "month": month, "days": (last - first).days + 1, "rows": rows,
            "departments": departments_of(users, names)}


# ----------------------------------------------------------------- files

def _dmy(d) -> str:
    return d.strftime("%d-%m-%Y") if d else ""


def _hm(dt) -> str:
    return timezone.localtime(dt).strftime("%H:%M") if dt else ""


def _hours(minutes) -> str:
    return f"{minutes // 60}:{minutes % 60:02d}" if minutes else ""


LABEL = {"present": "Present", "late": "Late", "half_day": "Half day", "absent": "Absent", "leave": "Leave",
         "holiday": "Holiday", "week_off": "Week off", "not_checked_in": "Not checked in"}
LETTER = {"present": "P", "late": "P", "half_day": "H", "absent": "A", "leave": "L", "holiday": "HO", "week_off": "WO"}
REPORT_TYPES = ("daily", "monthly_summary", "monthly_register", "late", "leave")


def _ident(u, names):
    return [u.employee_code, u.get_full_name() or u.username, names.get(u.department, u.department or ""), u.designation]


def build_report(kind, users, params):
    """(title, filename stem, header, rows) for one report type."""
    names = department_names()
    ident = ["Employee ID", "Name", "Department", "Designation"]
    if kind == "daily":
        day = parse_day(params.get("date") or timezone.localdate().isoformat())
        data = team_day(users, day)
        rows = []
        for u, r in zip(users, data["rows"]):
            rows.append(_ident(u, names) + [u.whatsapp_phone, LABEL.get(r["status"], "") if r["status"] else "",
                                            _hm(r["check_in"]), _hm(r["check_out"]), _hours(r["working_minutes"]),
                                            "Yes" if r["is_late"] else "", "Yes" if r["is_early_checkout"] else "", r["note"]])
        return (f"Daily attendance {_dmy(day)}", f"attendance-daily-{day.isoformat()}",
                ident + ["Phone", "Status", "Check in", "Check out", "Hours", "Late", "Early out", "Note"], rows)

    if kind in ("monthly_summary", "monthly_register"):
        year, month, first, last = parse_month(params)
        statuses, _ = month_statuses(users, first, last)
        stem = f"{year}-{month:02d}"
        if kind == "monthly_summary":
            rows = []
            for u in users:
                t = totals_for(statuses[u.id])
                tt = t["totals"]
                rows.append(_ident(u, names) + [t["present_days"], tt["late"], tt["half_day"], tt["absent"], tt["leave"],
                                                tt["week_off"], tt["holiday"], t["late_marks"], t["total_hours"]])
            return (f"Monthly attendance summary {first:%B %Y}", f"attendance-monthly-summary-{stem}",
                    ident + ["Present", "of which late", "Half day", "Absent", "Leave", "Week off", "Holiday", "Late marks", "Hours worked"],
                    rows)
        days = [first + timedelta(days=i) for i in range((last - first).days + 1)]
        rows = []
        for u in users:
            letters = [LETTER.get(statuses[u.id][d][0] or "", "") for d in days]
            count = {k: letters.count(k) for k in ("P", "A", "H", "L", "WO", "HO")}
            rows.append(_ident(u, names)[:3] + letters + [count[k] for k in ("P", "A", "H", "L", "WO", "HO")])
        return (f"Attendance register {first:%B %Y}", f"attendance-register-{stem}",
                ident[:3] + [f"{d.day:02d} {d:%a}" for d in days] + ["P", "A", "H", "L", "WO", "HO"], rows)

    if kind == "late":
        first, last = parse_range(params)
        start = services._parse_hhmm(services.cfg()["work_start"], time(9, 30))
        recs = (Attendance.objects.filter(user__in=[u.id for u in users], date__range=(first, last), is_late=True)
                .select_related("user").order_by("date", "check_in"))
        rows = []
        for a in recs:
            late_by = ""
            if a.check_in:
                delta = datetime.combine(a.date, timezone.localtime(a.check_in).time()) - datetime.combine(a.date, start)
                late_by = max(0, int(delta.total_seconds() // 60))
            rows.append([_dmy(a.date)] + _ident(a.user, names) + [_hm(a.check_in), late_by, _hm(a.check_out), _hours(a.working_minutes)])
        return (f"Late arrivals {_dmy(first)} to {_dmy(last)}", f"attendance-late-{first.isoformat()}-to-{last.isoformat()}",
                ["Date"] + ident + ["Check in", "Minutes late", "Check out", "Hours"], rows)

    if kind == "leave":
        first, last = parse_range(params)
        leaves = (LeaveRequest.objects.filter(user__in=[u.id for u in users], start_date__lte=last, end_date__gte=first)
                  .select_related("user", "leave_type", "reviewed_by").order_by("start_date", "user__first_name"))
        rows = [_ident(l.user, names) + [l.leave_type.name, _dmy(l.start_date), _dmy(l.end_date), l.days,
                                          l.get_status_display(), l.reason,
                                          (l.reviewed_by.get_full_name() or l.reviewed_by.username) if l.reviewed_by else ""]
                for l in leaves]
        return (f"Leave register {_dmy(first)} to {_dmy(last)}", f"leave-register-{first.isoformat()}-to-{last.isoformat()}",
                ident + ["Leave type", "From", "To", "Days", "Status", "Reason", "Reviewed by"], rows)

    raise ValidationError({"type": f"Use one of: {', '.join(REPORT_TYPES)}."})


def to_csv(header, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    from config.exports import safe_row
    w.writerow(safe_row(header))
    w.writerows(safe_row(r) for r in rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM: Excel opens it as UTF-8


def to_xlsx(title, header, rows, extra_sheet=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Report"
    from config.exports import safe_row
    ws.append(safe_row(header))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(safe_row(r))
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(4, width + 2))
    if extra_sheet:
        name, lines = extra_sheet
        notes = wb.create_sheet(name)
        for line in lines:
            notes.append([line])
        notes.column_dimensions["A"].width = 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------- import

TEMPLATE_HEADER = ["Employee ID", "Date (YYYY-MM-DD)", "Check in (HH:MM)", "Check out (HH:MM)", "Status", "Note"]
TEMPLATE_NOTES = [
    "One row per person per day. Columns: Employee ID, Date (YYYY-MM-DD), Check in (HH:MM), Check out (HH:MM), Status, Note.",
    "Employee ID is the person's employee ID in People (their username works too).",
    "Status: present, half_day, absent or leave. Leave it blank to work it out from the times, like a normal punch.",
    "Times are 24-hour (09:05, 18:30). Dates in the future are refused. An existing day is overwritten.",
]
STATUS_ALIASES = {
    "present": "present", "p": "present", "late": "present",
    "half_day": "half_day", "half day": "half_day", "halfday": "half_day", "h": "half_day", "hd": "half_day",
    "absent": "absent", "a": "absent",
    "leave": "leave", "l": "leave", "on leave": "leave",
}
COLUMN_KEYS = (("employee", "employee_id"), ("emp", "employee_id"), ("code", "employee_id"), ("date", "date"),
               ("check in", "check_in"), ("checkin", "check_in"), ("in", "check_in"),
               ("check out", "check_out"), ("checkout", "check_out"), ("out", "check_out"),
               ("status", "status"), ("note", "note"), ("remark", "note"))


def _column(name) -> str | None:
    text = re.sub(r"[_\-]+", " ", str(name or "").strip().lower())
    for prefix, key in COLUMN_KEYS:
        if text == prefix or text.startswith(prefix + " ") or text.startswith(prefix + "("):
            return key
    return None


def read_rows(upload):
    """[(line, {key: raw})] from an uploaded XLSX or CSV."""
    name = (getattr(upload, "name", "") or "").lower()
    if upload.size > IMPORT_MAX_BYTES:
        raise ValidationError({"file": "The file is larger than 2 MB."})
    raw = upload.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise ValidationError({"file": "This XLSX file could not be read."})
        rows_in = wb.worksheets[0].iter_rows(values_only=True, max_col=IMPORT_MAX_COLUMNS)
    elif name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        rows_in = (r[:IMPORT_MAX_COLUMNS] for r in csv.reader(io.StringIO(text)))
    else:
        raise ValidationError({"file": "Upload an .xlsx or .csv file."})
    # Read at most a bounded number of lines: a 2 MB XLSX can unpack to millions of rows (a "zip bomb").
    limit = IMPORT_MAX_ROWS * 2 + 2
    try:
        table = [list(r) for r in islice(rows_in, limit)]
    except Exception:
        raise ValidationError({"file": "This file could not be read."})
    if len(table) >= limit:
        raise ValidationError({"file": f"At most {IMPORT_MAX_ROWS} rows per file."})
    if not table:
        raise ValidationError({"file": "The file is empty."})
    keys = [_column(h) for h in table[0]]
    if "employee_id" not in keys or "date" not in keys:
        raise ValidationError({"file": "The first row must have the column names (Employee ID, Date, …). Download the template."})
    out = []
    for i, values in enumerate(table[1:], start=2):
        row = {}
        for key, value in zip(keys, values):
            if key and key not in row:
                row[key] = value
        if all(v is None or str(v).strip() == "" for v in row.values()):
            continue
        out.append((i, row))
    if len(out) > IMPORT_MAX_ROWS:
        raise ValidationError({"file": f"At most {IMPORT_MAX_ROWS} rows per file."})
    return out


def _cell_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _cell_time(value):
    """(time | None, ok)."""
    if value is None or str(value).strip() == "":
        return None, True
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0), True
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0), True
    if isinstance(value, float) and 0 <= value < 1:          # Excel stores a time as a fraction of a day
        minutes = round(value * 24 * 60)
        return time(minutes // 60 % 24, minutes % 60), True
    text = str(value).strip().upper()
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p", "%I:%M:%S %p"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0), True
        except ValueError:
            pass
    return None, False


def _code(value) -> str:
    """The employee ID as text (an Excel number cell 1001.0 -> "1001")."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value if value is not None else "").strip()


def _find_people(codes):
    """{lower code: user}: by employee code first, then by username (former staff included)."""
    wanted = {c.lower() for c in codes if c}
    if not wanted:
        return {}
    users = list(User.objects.only("id", "username", "first_name", "last_name", "employee_code", "is_active").order_by("id"))
    found = {}
    for u in users:
        if u.employee_code and u.employee_code.lower() in wanted:
            found.setdefault(u.employee_code.lower(), u)
    for u in users:
        if u.username.lower() in wanted:
            found.setdefault(u.username.lower(), u)
    return found


def validate_rows(parsed):
    """Rows for the preview: each with ok/error, and what will be stored."""
    today = timezone.localdate()
    people = _find_people([_code(r.get("employee_id")) for _, r in parsed])
    seen = set()
    out = []
    for line, r in parsed:
        code = _code(r.get("employee_id"))
        user = people.get(code.lower()) if code else None
        day = _cell_date(r.get("date"))
        t_in, in_ok = _cell_time(r.get("check_in"))
        t_out, out_ok = _cell_time(r.get("check_out"))
        status_raw = str(r.get("status") or "").strip().lower()
        status = STATUS_ALIASES.get(status_raw) if status_raw else ""
        note = str(r.get("note") or "").strip()[:200]
        error = ""
        if not code:
            error = "Employee ID is missing."
        elif not user:
            error = f"No employee with ID '{code}'."
        elif not day:
            error = "Date is not a date (use YYYY-MM-DD)."
        elif day > today:
            error = "Date is in the future."
        elif not in_ok:
            error = "Check in is not a time (use HH:MM)."
        elif not out_ok:
            error = "Check out is not a time (use HH:MM)."
        elif t_out and not t_in:
            error = "Check out without a check in."
        elif t_in and t_out and t_out < t_in:
            error = "Check out is before check in."
        elif status is None:
            error = "Status must be present, half_day, absent or leave."
        elif status in ("absent", "leave") and (t_in or t_out):
            error = f"A day marked {status} cannot have check-in or check-out times."
        elif not status and not t_in:
            error = "Give a status or a check-in time."
        elif (user.id, day) in seen:
            error = "This person and date are already in the file."
        if not error:
            seen.add((user.id, day))
        out.append({
            "line": line, "employee_id": code, "user_id": user.id if user else None,
            "name": (user.get_full_name() or user.username) if user else "",
            "date": day.isoformat() if day else str(r.get("date") or ""),
            "status": status or ("present" if t_in else ""),
            "check_in": t_in.strftime("%H:%M") if t_in else "",
            "check_out": t_out.strftime("%H:%M") if t_out else "",
            "note": note, "ok": not error, "error": error,
            "_user": user, "_day": day, "_in": t_in, "_out": t_out, "_explicit": bool(status),
        })
    return out


def _apply_row(row, by_name, today_text, existing):
    conf = services.cfg()
    user, day = row["_user"], row["_day"]
    rec = existing.get((user.id, day)) or Attendance(user=user, date=day)
    tz = timezone.get_current_timezone()
    rec.check_in = timezone.make_aware(datetime.combine(day, row["_in"]), tz) if row["_in"] else None
    rec.check_out = timezone.make_aware(datetime.combine(day, row["_out"]), tz) if row["_out"] else None
    start = services._parse_hhmm(conf["work_start"], time(9, 30))
    end = services._parse_hhmm(conf["work_end"], time(18, 30))
    late_cutoff = (datetime.combine(day, start) + timedelta(minutes=conf["grace_minutes"])).time()
    rec.is_late = bool(row["_in"] and row["_in"] > late_cutoff)
    rec.is_early_checkout = bool(row["_out"] and row["_out"] < end)
    rec.working_minutes = (int((rec.check_out - rec.check_in).total_seconds() // 60)
                           if rec.check_in and rec.check_out else 0)
    status = row["status"]
    if status == "present":
        if not row["_explicit"] and rec.check_out and rec.working_minutes / 60 < conf["half_day_hours"]:
            status = "half_day"            # blank status: decided like a normal punch
        else:
            status = "late" if rec.is_late else "present"
    rec.status = status
    if status != "leave":
        rec.leave_request = None
    rec.note = (f"Imported by {by_name} on {today_text}" + (f": {row['note']}" if row["note"] else ""))[:300]
    rec.save()


def import_attendance(upload, by_user, dry_run=True) -> dict:
    rows = validate_rows(read_rows(upload))
    good = [r for r in rows if r["ok"]]
    existing = {}
    if good:
        days = {r["_day"] for r in good}
        for a in Attendance.objects.filter(user__in={r["_user"].id for r in good}, date__in=days):
            existing[(a.user_id, a.date)] = a
    created = sum(1 for r in good if (r["_user"].id, r["_day"]) not in existing)
    updated = len(good) - created
    if not dry_run and good:
        by_name = by_user.get_full_name() or by_user.username
        today_text = _dmy(timezone.localdate())
        with transaction.atomic():
            for r in good:
                _apply_row(r, by_name, today_text, existing)
    skipped = [{"line": r["line"], "error": r["error"]} for r in rows if not r["ok"]]
    public = [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows]
    return {"dry_run": dry_run, "rows": public, "created": created, "updated": updated,
            "skipped": len(skipped), "skipped_rows": skipped}
