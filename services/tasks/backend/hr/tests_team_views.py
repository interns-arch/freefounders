"""Attendance for managers (PLAN.md §19.15): team day / month, report files and the attendance import."""
import csv
import io
from datetime import date, datetime, time, timedelta

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from accounts.models import Role, User
from crm.models import Holiday
from hr import services
from hr.models import Attendance, AttendanceStatus, LeaveRequest, LeaveStatus, LeaveType
from tenancy.company_settings import forget

# A past month with known weekdays: September 2026 starts on a Tuesday; Sundays (week-off "6") are 6, 13, 20, 27.
SEP = (2026, 9)
MON = date(2026, 9, 14)
SUN = date(2026, 9, 13)
HOLI = date(2026, 9, 15)


def at(day, hh, mm):
    return timezone.make_aware(datetime.combine(day, time(hh, mm)))


class _Base(TestCase):
    def setUp(self):
        forget()
        mk = User.objects.create_user
        self.hr = mk("hr", "hr@x.com", "pass@12345", role=Role.HR_MANAGER, department="hr", first_name="Hema")
        self.manager = mk("sm", "sm@x.com", "pass@12345", role=Role.SALES_MANAGER, department="sales", first_name="Sam")
        self.rahul = mk("rahul", "rahul@x.com", "pass@12345", role=Role.SALES_EXECUTIVE, department="sales",
                        first_name="Rahul", last_name="Verma", employee_code="CT-001", whatsapp_phone="919800000001",
                        designation="Sales executive")
        self.priya = mk("priya", "priya@x.com", "pass@12345", role=Role.SALES_EXECUTIVE, department="purchase",
                        first_name="Priya", employee_code="CT-002")
        self.gone = mk("gone", "gone@x.com", "pass@12345", role=Role.SALES_EXECUTIVE, department="sales",
                       first_name="Gone", employee_code="CT-009", is_active=False)
        Holiday.objects.create(name="Festival", date=HOLI)

    def tearDown(self):
        forget()

    def as_(self, user):
        c = APIClient()
        c.force_authenticate(user)
        return c

    def ids(self, res):
        self.assertEqual(res.status_code, 200, res.content)
        return {r["user_id"] for r in res.json()["rows"]}


class ScopeTests(_Base):
    def test_manager_sees_only_their_department(self):
        ids = self.ids(self.as_(self.manager).get("/api/attendance/team_day/"))
        self.assertIn(self.rahul.id, ids)
        self.assertNotIn(self.priya.id, ids)
        self.assertNotIn(self.hr.id, ids)

    def test_hr_sees_the_company_and_inactive_on_request(self):
        ids = self.ids(self.as_(self.hr).get("/api/attendance/team_day/"))
        self.assertTrue({self.rahul.id, self.priya.id, self.manager.id, self.hr.id} <= ids)
        self.assertNotIn(self.gone.id, ids)
        self.assertIn(self.gone.id, self.ids(self.as_(self.hr).get("/api/attendance/team_month/", {"inactive": "1"})))

    def test_member_gets_403_everywhere(self):
        c = self.as_(self.rahul)
        for url in ("/api/attendance/team_day/", "/api/attendance/team_month/", "/api/attendance/report/?type=daily"):
            self.assertEqual(c.get(url).status_code, 403, url)
        self.assertEqual(c.get("/api/attendance/import_template/").status_code, 403)
        self.assertEqual(c.post("/api/attendance/import/", {}).status_code, 403)

    def test_manager_cannot_import(self):
        self.assertEqual(self.as_(self.manager).get("/api/attendance/import_template/").status_code, 403)
        self.assertEqual(self.as_(self.manager).post("/api/attendance/import/", {}).status_code, 403)

    def test_filters(self):
        c = self.as_(self.hr)
        self.assertEqual(self.ids(c.get("/api/attendance/team_day/", {"department": "purchase"})), {self.priya.id})
        self.assertEqual(self.ids(c.get("/api/attendance/team_day/", {"q": "ct-001"})), {self.rahul.id})
        self.assertEqual(self.ids(c.get("/api/attendance/team_day/", {"q": "98000"})), {self.rahul.id})
        self.assertEqual(self.ids(c.get("/api/attendance/team_day/", {"q": "verma"})), {self.rahul.id})
        res = c.get("/api/attendance/team_day/")
        self.assertIn({"code": "purchase", "name": "Purchase"}, res.json()["departments"])


class DayTests(_Base):
    def status_on(self, day, user=None):
        res = self.as_(self.hr).get("/api/attendance/team_day/", {"date": day.isoformat()})
        self.assertEqual(res.status_code, 200, res.content)
        return next(r for r in res.json()["rows"] if r["user_id"] == (user or self.rahul).id)

    def test_classification(self):
        self.assertEqual(self.status_on(SUN)["status"], "week_off")
        self.assertEqual(self.status_on(HOLI)["status"], "holiday")
        self.assertEqual(self.status_on(MON)["status"], "absent")
        later = timezone.localdate() + timedelta(days=1)
        while services.is_week_off(later):
            later += timedelta(days=1)
        self.assertIsNone(self.status_on(later)["status"])
        today = timezone.localdate()
        if not services.is_week_off(today) and not services.is_holiday(today):
            self.assertEqual(self.status_on(today)["status"], "not_checked_in")

    def test_leave_and_punches(self):
        lt = LeaveType.objects.create(name="Casual")
        leave = LeaveRequest.objects.create(user=self.rahul, leave_type=lt, start_date=MON, end_date=MON, days=1,
                                            reason="x", status=LeaveStatus.APPROVED)
        services.apply_leave_to_attendance(leave)
        self.assertEqual(self.status_on(MON)["status"], "leave")
        Attendance.objects.create(user=self.priya, date=MON, check_in=at(MON, 10, 5), check_out=at(MON, 19, 0),
                                  status="late", is_late=True, working_minutes=535)
        row = self.status_on(MON, self.priya)
        self.assertEqual((row["status"], row["is_late"], row["working_minutes"]), ("late", True, 535))
        self.assertEqual(row["employee_code"], "CT-002")

    def test_bad_date(self):
        self.assertEqual(self.as_(self.hr).get("/api/attendance/team_day/", {"date": "10-09-2026"}).status_code, 400)


class MonthTests(_Base):
    def test_totals_match_monthly_report(self):
        Attendance.objects.create(user=self.rahul, date=date(2026, 9, 1), check_in=at(date(2026, 9, 1), 9, 20),
                                  check_out=at(date(2026, 9, 1), 18, 40), status="present", working_minutes=560)
        Attendance.objects.create(user=self.rahul, date=date(2026, 9, 2), check_in=at(date(2026, 9, 2), 10, 0),
                                  check_out=at(date(2026, 9, 2), 19, 0), status="late", is_late=True, working_minutes=540)
        Attendance.objects.create(user=self.rahul, date=date(2026, 9, 3), check_in=at(date(2026, 9, 3), 9, 0),
                                  check_out=at(date(2026, 9, 3), 11, 0), status="half_day", working_minutes=120)
        Attendance.objects.create(user=self.rahul, date=date(2026, 9, 4), status="leave")
        res = self.as_(self.hr).get("/api/attendance/team_month/", {"year": SEP[0], "month": SEP[1]})
        self.assertEqual(res.status_code, 200, res.content)
        row = next(r for r in res.json()["rows"] if r["user_id"] == self.rahul.id)
        expected = services.monthly_report(self.rahul, *SEP)
        self.assertEqual(row["totals"], expected["totals"])
        self.assertEqual(row["total_hours"], expected["total_hours"])
        self.assertEqual(row["present_days"], 2)
        self.assertEqual(row["late_marks"], 1)
        self.assertEqual(row["totals"]["holiday"], 1)
        self.assertEqual(row["totals"]["week_off"], 4)

    def test_no_query_per_person(self):
        for i in range(8):
            User.objects.create_user(f"u{i}", f"u{i}@x.com", "pass@12345", department="sales")
        c = self.as_(self.hr)
        c.get("/api/attendance/team_month/", {"year": SEP[0], "month": SEP[1]})        # warm caches
        with self.assertNumQueriesLessThan(25):
            c.get("/api/attendance/team_month/", {"year": SEP[0], "month": SEP[1]})

    def assertNumQueriesLessThan(self, n):
        test = self

        class _Ctx:
            def __enter__(self):
                from django.db import connection
                from django.test.utils import CaptureQueriesContext
                self.cap = CaptureQueriesContext(connection)
                self.cap.__enter__()

            def __exit__(self, *exc):
                self.cap.__exit__(*exc)
                test.assertLess(len(self.cap.captured_queries), n, [q["sql"] for q in self.cap.captured_queries])
        return _Ctx()


class ReportTests(_Base):
    def setUp(self):
        super().setUp()
        Attendance.objects.create(user=self.rahul, date=MON, check_in=at(MON, 10, 0), check_out=at(MON, 19, 0),
                                  status="late", is_late=True, working_minutes=540)

    def get(self, **params):
        return self.as_(self.hr).get("/api/attendance/report/", params)

    def xlsx(self, res):
        self.assertEqual(res.status_code, 200, res.content[:300])
        ws = load_workbook(io.BytesIO(res.content)).active
        return [list(r) for r in ws.iter_rows(values_only=True)], ws

    def test_daily_xlsx(self):
        res = self.get(type="daily", date=MON.isoformat(), format="xlsx")
        self.assertIn('filename="attendance-daily-2026-09-14.xlsx"', res["Content-Disposition"])
        rows, ws = self.xlsx(res)
        self.assertEqual(rows[0][:6], ["Employee ID", "Name", "Department", "Designation", "Phone", "Status"])
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")
        rahul = next(r for r in rows if r[0] == "CT-001")
        self.assertEqual(rahul[1:], ["Rahul Verma", "Sales", "Sales executive", "919800000001", "Late", "10:00", "19:00",
                                     "9:00", "Yes", None, None])
        priya = next(r for r in rows if r[0] == "CT-002")
        self.assertEqual(priya[5], "Absent")

    def test_daily_csv(self):
        res = self.get(type="daily", date=MON.isoformat(), format="csv")
        self.assertEqual(res.status_code, 200)
        self.assertIn('attendance-daily-2026-09-14.csv', res["Content-Disposition"])
        rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0][0], "Employee ID")
        self.assertIn(["CT-001", "Rahul Verma", "Sales", "Sales executive", "919800000001", "Late", "10:00", "19:00", "9:00",
                       "Yes", "", ""], rows)

    def test_monthly_summary_and_register(self):
        rows, _ = self.xlsx(self.get(type="monthly_summary", year=2026, month=9, format="xlsx"))
        self.assertEqual(rows[0][4:6], ["Present", "of which late"])
        rahul = next(r for r in rows if r[0] == "CT-001")
        self.assertEqual(rahul[4], 1)
        rows, _ = self.xlsx(self.get(type="monthly_register", year=2026, month=9, format="xlsx"))
        self.assertEqual(rows[0][3], "01 Tue")
        self.assertEqual(len(rows[0]), 3 + 30 + 6)
        rahul = next(r for r in rows if r[0] == "CT-001")
        self.assertEqual(rahul[3 + 13], "P")      # 14 Sep (late = present)
        self.assertEqual(rahul[3 + 12], "WO")     # 13 Sep, Sunday
        self.assertEqual(rahul[3 + 14], "HO")     # 15 Sep, holiday
        self.assertEqual(rahul[3], "A")
        self.assertEqual(rahul[-6:-4], [1, 30 - 4 - 1 - 1])

    def test_late_and_leave_registers(self):
        res = self.get(type="late", **{"from": "2026-09-01", "to": "2026-09-30"}, format="csv")
        rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))
        self.assertEqual(rows[1][:2], ["14-09-2026", "CT-001"])
        self.assertEqual(rows[1][6], "30")        # 10:00 against a 09:30 start
        lt = LeaveType.objects.create(name="Casual")
        LeaveRequest.objects.create(user=self.priya, leave_type=lt, start_date=MON, end_date=MON, days=1, reason="Trip")
        res = self.get(type="leave", **{"from": "2026-09-01", "to": "2026-09-30"}, format="xlsx")
        self.assertIn("leave-register-2026-09-01-to-2026-09-30.xlsx", res["Content-Disposition"])
        rows, _ = self.xlsx(res)
        self.assertEqual(rows[1][:7], ["CT-002", "Priya", "Purchase", None, "Casual", "14-09-2026", "14-09-2026"])

    def test_bad_params(self):
        self.assertEqual(self.get(type="nope").status_code, 400)
        self.assertEqual(self.get(type="daily", format="pdf").status_code, 400)
        self.assertEqual(self.get(type="late", **{"from": "2026-09-10", "to": "2026-09-01"}).status_code, 400)

    def test_manager_report_is_their_team_only(self):
        res = self.as_(self.manager).get("/api/attendance/report/", {"type": "daily", "date": MON.isoformat(), "format": "csv"})
        text = res.content.decode("utf-8-sig")
        self.assertIn("CT-001", text)
        self.assertNotIn("CT-002", text)


class ImportTests(_Base):
    HEADER = ["Employee ID", "Date (YYYY-MM-DD)", "Check in (HH:MM)", "Check out (HH:MM)", "Status", "Note"]

    def csv_file(self, rows):
        buf = io.StringIO()
        csv.writer(buf).writerows([self.HEADER, *rows])
        return SimpleUploadedFile("days.csv", buf.getvalue().encode(), content_type="text/csv")

    def post(self, f, dry_run):
        res = self.as_(self.hr).post("/api/attendance/import/", {"file": f, "dry_run": "true" if dry_run else "false"},
                                     format="multipart")
        self.assertEqual(res.status_code, 200, res.content)
        return res.json()

    def test_template(self):
        res = self.as_(self.hr).get("/api/attendance/import_template/", {"format": "xlsx"})
        wb = load_workbook(io.BytesIO(res.content))
        self.assertEqual([c.value for c in wb.active[1]], self.HEADER)
        self.assertEqual(wb.sheetnames[1], "How to fill")
        res = self.as_(self.hr).get("/api/attendance/import_template/", {"format": "csv"})
        self.assertEqual(res.content.decode("utf-8-sig").strip(), ",".join(self.HEADER))

    def test_dry_run_validation(self):
        future = (timezone.localdate() + timedelta(days=2)).isoformat()
        out = self.post(self.csv_file([
            ["CT-001", "2026-09-14", "09:20", "18:40", "", ""],
            ["NOPE", "2026-09-14", "09:20", "18:40", "", ""],
            ["CT-002", "14/13/2026", "09:20", "", "", ""],
            ["CT-002", future, "09:20", "", "", ""],
            ["CT-002", "2026-09-14", "18:00", "09:00", "", ""],
            ["CT-002", "2026-09-16", "", "", "maybe", ""],
            ["CT-001", "2026-09-14", "", "", "absent", ""],
            ["rahul", "2026-09-17", "", "", "absent", ""],
        ]), dry_run=True)
        errors = [r["error"] for r in out["rows"]]
        self.assertEqual(errors[0], "")
        self.assertIn("No employee", errors[1])
        self.assertIn("not a date", errors[2])
        self.assertIn("future", errors[3])
        self.assertIn("before check in", errors[4])
        self.assertIn("Status must be", errors[5])
        self.assertIn("already in the file", errors[6])
        self.assertEqual(errors[7], "", "username works as a fallback")
        self.assertEqual((out["created"], out["updated"], out["skipped"]), (2, 0, 6))
        self.assertFalse(Attendance.objects.exists(), "a dry run stores nothing")

    def test_apply_and_reimport_updates(self):
        Attendance.objects.create(user=self.priya, date=MON, status="absent")
        rows = [
            ["CT-001", "2026-09-14", "10:00", "19:00", "", "from SalaryBox"],
            ["CT-001", "2026-09-16", "09:00", "11:00", "", ""],
            ["CT-002", "2026-09-14", "09:15", "18:45", "present", ""],
            ["CT-009", "2026-09-16", "", "", "leave", ""],
            ["CT-002", "2026-09-17", "", "", "half_day", ""],
        ]
        out = self.post(self.csv_file(rows), dry_run=False)
        self.assertEqual((out["created"], out["updated"], out["skipped"]), (4, 1, 0))
        a = Attendance.objects.get(user=self.rahul, date=MON)
        self.assertEqual((a.status, a.is_late, a.working_minutes, a.is_early_checkout), ("late", True, 540, False))
        self.assertEqual(timezone.localtime(a.check_in).time(), time(10, 0))
        self.assertTrue(a.note.startswith("Imported by Hema on "), a.note)
        self.assertTrue(a.note.endswith(": from SalaryBox"))
        short = Attendance.objects.get(user=self.rahul, date=date(2026, 9, 16))
        self.assertEqual((short.status, short.working_minutes, short.is_early_checkout), ("half_day", 120, True))
        p = Attendance.objects.get(user=self.priya, date=MON)
        self.assertEqual((p.status, p.working_minutes), ("present", 570))
        self.assertEqual(Attendance.objects.get(user=self.gone, date=date(2026, 9, 16)).status, "leave")
        self.assertEqual(Attendance.objects.get(user=self.priya, date=date(2026, 9, 17)).status, "half_day")

        rows[0] = ["CT-001", "2026-09-14", "09:10", "18:40", "", ""]
        out = self.post(self.csv_file(rows), dry_run=False)
        self.assertEqual((out["created"], out["updated"]), (0, 5))
        self.assertEqual(Attendance.objects.filter(user=self.rahul, date=MON).count(), 1)
        a.refresh_from_db()
        self.assertEqual((a.status, a.is_late, a.working_minutes), ("present", False, 570))

    def test_xlsx_upload_with_excel_cells(self):
        wb = Workbook()
        ws = wb.active
        ws.append(self.HEADER)
        ws.append(["CT-001", datetime(2026, 9, 14), time(9, 25), time(18, 35), "Present", None])
        ws.append([None, None, None, None, None, None])
        buf = io.BytesIO()
        wb.save(buf)
        f = SimpleUploadedFile("days.xlsx", buf.getvalue())
        out = self.post(f, dry_run=False)
        self.assertEqual((out["created"], out["skipped"]), (1, 0), out)
        self.assertEqual(Attendance.objects.get(user=self.rahul, date=MON).working_minutes, 550)

    def test_bad_files(self):
        res = self.as_(self.hr).post("/api/attendance/import/", {"file": SimpleUploadedFile("a.pdf", b"x")}, format="multipart")
        self.assertEqual(res.status_code, 400)
        f = SimpleUploadedFile("a.csv", b"Name,Day\nx,y\n")
        res = self.as_(self.hr).post("/api/attendance/import/", {"file": f}, format="multipart")
        self.assertEqual(res.status_code, 400)
        self.assertIn("column names", str(res.json()))


class ManualMarkDateTests(_Base):
    def test_past_date_and_future_refused(self):
        c = self.as_(self.hr)
        res = c.post("/api/attendance/manual_mark/", {"user": self.rahul.id, "status": "half_day", "date": MON.isoformat(),
                                                      "reason": "forgot"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(Attendance.objects.get(user=self.rahul, date=MON).status, AttendanceStatus.HALF_DAY)
        future = (timezone.localdate() + timedelta(days=1)).isoformat()
        self.assertEqual(c.post("/api/attendance/manual_mark/", {"user": self.rahul.id, "date": future}, format="json").status_code, 400)
