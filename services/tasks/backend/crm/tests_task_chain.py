"""Sub-tasks and the Task chain: whoever holds a task gives parts of it to
others (who can pass parts on, up to MAX_DEPTH levels); a part's completion
goes to the person who gave it; the task above can't be completed until
every part is done and accepted; effort is split, never earned twice; and
admins see the whole chain. Ported from the original app (1dabd35) with the
Kaampilot additions at the end (PLAN.md §19.41): passing on a part needs no
Assign permission, a self-made task does; categories of the new notifications;
change requests keep the old due date for the ranking and the Excel report."""
from datetime import timedelta
from io import BytesIO
from unittest import mock

from django.utils import timezone
from openpyxl import load_workbook

from notifications.models import Notification
from .chain import MAX_DEPTH
from .models import ChangeRequestStatus, Task, TaskChangeRequest, TaskCompletion, TaskStatus
from .tests_task_engine import Base, make
from accounts.models import Role


class ChainBase(Base):
    def setUp(self):
        super().setUp()
        self.now = timezone.now()
        # manager (A) gives rahul (B) the main task
        self.main = Task.objects.create(
            title="Big job", assigned_to=self.rahul, created_by=self.manager,
            effort_minutes=240, due_at=self.now + timedelta(days=5))

    # ---- helpers ---------------------------------------------------------
    def give(self, as_user, parent, to, *, effort=60, due_in_days=2, **extra):
        self.as_(as_user)
        return self.client.post("/api/tasks/", {
            "title": f"Part for {to.username}", "assigned_to": to.id,
            "parent": parent.id, "effort_minutes": effort,
            "due_at": (self.now + timedelta(days=due_in_days)).isoformat(), **extra,
        }, format="json")

    def complete(self, as_user, task, minutes=30):
        self.as_(as_user)
        return self.client.post(f"/api/tasks/{task.id}/complete/",
                                {"remarks": "done it", "actual_minutes": minutes})

    def accept(self, as_user, task):
        c = TaskCompletion.objects.filter(task=task).latest("created_at")
        self.as_(as_user)
        return self.client.post(f"/api/task-completions/{c.id}/review/",
                                {"decision": "approved"}, format="json")

    def part(self, res):
        self.assertEqual(res.status_code, 201, res.data)
        return Task.objects.get(pk=res.data["id"])


class ChainTests(ChainBase):
    # ---- giving parts ------------------------------------------------------
    def test_holder_gives_parts_to_several_people(self):
        a = self.part(self.give(self.rahul, self.main, self.amit, frequency="daily"))
        b = self.part(self.give(self.rahul, self.main, self.vikram))
        for p in (a, b):
            self.assertEqual(p.parent_id, self.main.id)
            self.assertEqual(p.created_by_id, self.rahul.id)
            self.assertEqual(p.frequency, "one_time")     # parts never repeat
        self.assertTrue(Notification.objects.filter(user=self.amit).exists())

    def test_only_the_holder_or_an_admin_can_give_parts(self):
        self.assertEqual(self.give(self.manager, self.main, self.amit).status_code, 403)
        self.assertEqual(self.give(self.amit, self.main, self.vikram).status_code, 400)  # can't see it
        self.part(self.give(self.admin, self.main, self.amit))

    def test_part_rules_date_effort_and_people(self):
        self.assertIn("due_at", self.give(self.rahul, self.main, self.amit, due_in_days=9).data)
        self.assertIn("effort_minutes", self.give(self.rahul, self.main, self.amit, effort=300).data)
        self.assertIn("assigned_to", self.give(self.rahul, self.main, self.rahul).data)
        self.part(self.give(self.rahul, self.main, self.amit, effort=200))
        # only 40 of 240 left to give now
        res = self.give(self.rahul, self.main, self.vikram, effort=60)
        self.assertEqual(res.status_code, 400)
        self.assertIn("40", str(res.data["effort_minutes"]))

    def test_parts_of_parts_up_to_the_limit(self):
        people = [make(f"p{i}", Role.SALES_EXECUTIVE) for i in range(MAX_DEPTH + 1)]
        holder, task = self.rahul, self.main
        for i in range(MAX_DEPTH):
            task = self.part(self.give(holder, task, people[i], effort=10))
            holder = people[i]
        res = self.give(holder, task, people[MAX_DEPTH], effort=5)
        self.assertEqual(res.status_code, 400)
        self.assertIn("levels", str(res.data))

    def test_cannot_give_parts_of_a_finished_task(self):
        Task.objects.filter(pk=self.main.pk).update(status=TaskStatus.DONE)
        self.assertEqual(self.give(self.rahul, self.main, self.amit).status_code, 400)

    # ---- completing ----------------------------------------------------------
    def test_main_task_waits_for_every_part_to_be_done_and_accepted(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        res = self.complete(self.rahul, self.main)
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data["needs"], "subtasks")
        self.assertIn(p.code, str(res.data["detail"]))

        self.assertEqual(self.complete(self.amit, p).status_code, 200)
        review = TaskCompletion.objects.get(task=p)
        self.assertEqual(review.approver_id, self.rahul.id)  # goes to who gave the part
        res = self.complete(self.rahul, self.main)              # done, not accepted yet
        self.assertEqual(res.status_code, 400)
        self.assertIn("accept", str(res.data["detail"]))

        self.assertEqual(self.accept(self.rahul, p).status_code, 200)
        self.assertTrue(Notification.objects.filter(
            user=self.rahul, type="subtasks_all_done").exists())
        self.assertEqual(self.complete(self.rahul, self.main).status_code, 200)
        # and the main task's own completion goes to the original giver
        self.assertEqual(TaskCompletion.objects.get(task=self.main).approver_id, self.manager.id)

    def test_status_patch_is_blocked_too(self):
        self.part(self.give(self.rahul, self.main, self.amit))
        self.as_(self.rahul)
        res = self.client.patch(f"/api/tasks/{self.main.id}/", {"status": "done"}, format="json")
        self.assertEqual(res.status_code, 400)

    def test_sent_back_part_blocks_again(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        self.complete(self.amit, p)
        c = TaskCompletion.objects.get(task=p)
        self.as_(self.rahul)
        self.client.post(f"/api/task-completions/{c.id}/review/",
                         {"decision": "rejected", "remarks": "photo is missing here"},
                         format="json")
        self.assertEqual(Task.objects.get(pk=p.pk).status, TaskStatus.OPEN)
        self.assertEqual(self.complete(self.rahul, self.main).status_code, 400)

    def test_bulk_accept_of_last_part_tells_the_holder(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        self.complete(self.amit, p)
        c = TaskCompletion.objects.get(task=p)
        self.as_(self.rahul)
        with mock.patch("crm.task_views._in_background", side_effect=lambda fn: fn()):
            res = self.client.post("/api/task-completions/bulk_accept/", {"ids": [c.id]},
                                   format="json")
        self.assertEqual(res.data["accepted"], 1)
        self.assertTrue(Notification.objects.filter(
            user=self.rahul, type="subtasks_all_done").exists())

    # ---- cancel / due date ---------------------------------------------------
    def test_deleting_the_main_task_cancels_open_parts_and_restore_brings_them_back(self):
        open_part = self.part(self.give(self.rahul, self.main, self.amit))
        deeper = self.part(self.give(self.amit, open_part, self.vikram, effort=10))
        done_part = self.part(self.give(self.rahul, self.main, self.hr))
        Task.objects.filter(pk=done_part.pk).update(status=TaskStatus.DONE)
        self.as_(self.manager)                                # the giver may delete
        self.assertEqual(self.client.delete(f"/api/tasks/{self.main.id}/").status_code, 204)
        self.assertIsNotNone(Task.objects.get(pk=open_part.pk).deleted_at)
        self.assertIsNotNone(Task.objects.get(pk=deeper.pk).deleted_at)
        self.assertIsNone(Task.objects.get(pk=done_part.pk).deleted_at)   # finished work stays
        self.assertTrue(Notification.objects.filter(user=self.vikram, type="task_cancelled").exists())

        self.as_(self.admin)
        self.client.post(f"/api/tasks/{self.main.id}/restore/")
        self.assertIsNone(Task.objects.get(pk=open_part.pk).deleted_at)
        self.assertIsNone(Task.objects.get(pk=deeper.pk).deleted_at)

    def test_earlier_due_date_flags_parts_due_after_it(self):
        p = self.part(self.give(self.rahul, self.main, self.amit, due_in_days=4))
        self.as_(self.admin)
        res = self.client.patch(f"/api/tasks/{self.main.id}/", {
            "due_at": (self.now + timedelta(days=3)).isoformat()}, format="json")
        self.assertEqual(res.status_code, 200, res.data)
        n = Notification.objects.filter(user=self.rahul, type="subtask_due_after_parent")
        self.assertEqual(n.count(), 1)
        self.assertIn(p.code, n.first().title)

    # ---- effort is split, never earned twice -----------------------------------
    def test_effort_split_in_the_score(self):
        p = self.part(self.give(self.rahul, self.main, self.amit, effort=60))
        self.complete(self.amit, p)
        self.accept(self.rahul, p)
        self.complete(self.rahul, self.main)
        self.as_(self.admin)
        rows = {r["user"]: r for r in self.client.get(
            "/api/tasks/employees_report/?range=all").data["rows"]}
        self.assertEqual(rows[self.rahul.id]["time_assigned_minutes"], 180)
        self.assertEqual(rows[self.rahul.id]["time_earned_minutes"], 180)
        self.assertEqual(rows[self.amit.id]["time_earned_minutes"], 60)
        time = {r["user"]: r for r in self.client.get(
            "/api/tasks/time_report/?range=all").data["rows"]} \
            if self.client.get("/api/tasks/time_report/?range=all").status_code == 200 else None
        if time:
            self.assertEqual(time[self.rahul.id]["time_earned_minutes"], 180)

    # ---- what each person sees ---------------------------------------------------
    def test_detail_shows_parts_and_the_task_above(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        self.as_(self.rahul)
        d = self.client.get(f"/api/tasks/{self.main.id}/").data
        self.assertEqual([s["code"] for s in d["subtasks"]], [p.code])
        self.assertEqual(d["chain"]["effort_left"], 180)
        self.assertTrue(d["chain"]["can_give_part"])
        self.assertFalse(d["chain"]["can_view_chain"])           # not an admin
        self.as_(self.amit)
        d = self.client.get(f"/api/tasks/{p.id}/").data
        self.assertEqual(d["chain"]["parent"]["code"], self.main.code)
        self.assertEqual(d["chain"]["parent"]["assignee"],
                         self.rahul.get_full_name() or self.rahul.username)
        self.assertEqual(self.client.get(f"/api/tasks/{p.id}/chain/").status_code, 403)

    def test_admin_sees_the_whole_chain(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        deeper = self.part(self.give(self.amit, p, self.vikram, effort=10))
        self.as_(self.admin)
        d = self.client.get(f"/api/tasks/{deeper.id}/chain/").data
        t = d["tree"]
        self.assertEqual(t["code"], self.main.code)
        self.assertEqual(t["parts"][0]["code"], p.code)
        self.assertEqual(t["parts"][0]["parts"][0]["code"], deeper.code)
        self.assertEqual(t["parts"][0]["parts"][0]["level"], 2)
        self.assertEqual(d["summary"]["parts"], 2)
        self.assertEqual(d["summary"]["people"], 3)

        rows = self.client.get("/api/tasks/chains/").data["rows"]
        self.assertEqual([r["code"] for r in rows], [self.main.code])
        self.assertEqual(rows[0]["parts"], 2)
        self.as_(self.manager)
        self.assertEqual(self.client.get("/api/tasks/chains/").status_code, 403)

    def test_excel_report_shows_the_chain(self):
        self.part(self.give(self.rahul, self.main, self.amit))
        self.as_(self.admin)
        res = self.client.get(f"/api/tasks/person_export/?user={self.amit.id}&range=all")
        ws = load_workbook(BytesIO(res.content))["Tasks received"]
        rows = list(ws.iter_rows(values_only=True))
        head = rows[0]
        self.assertEqual(rows[1][head.index("Sub-task of")], self.main.code)
        self.assertEqual(rows[1][head.index("Main task")], self.main.code)


class ChainKaampilotTests(ChainBase):
    """What Kaampilot adds around the original rules (PLAN.md §19.41)."""

    def test_a_self_made_task_is_split_only_with_the_assign_permission(self):
        mine = Task.objects.create(title="My own", assigned_to=self.rahul, created_by=self.rahul,
                                   effort_minutes=60, due_at=self.now + timedelta(days=3))
        self.assertEqual(self.give(self.rahul, mine, self.amit).status_code, 403)
        self.as_(self.rahul)
        d = self.client.get(f"/api/tasks/{mine.id}/").data
        self.assertFalse(d["chain"]["can_give_part"])
        # a manager (Assign permission) may split their own task
        theirs = Task.objects.create(title="Mgr own", assigned_to=self.manager, created_by=self.manager,
                                     effort_minutes=60, due_at=self.now + timedelta(days=3))
        self.part(self.give(self.manager, theirs, self.amit))

    def test_people_who_can_get_a_part(self):
        self.as_(self.rahul)
        ids = {p["id"] for p in self.client.get(f"/api/tasks/assignees/?part_of={self.main.id}").data}
        self.assertIn(self.amit.id, ids)
        self.assertNotIn(self.rahul.id, ids)              # never the holder
        self.assertEqual({p["id"] for p in self.client.get("/api/tasks/assignees/").data}, {self.rahul.id})
        self.as_(self.vikram)                                # cannot see the task
        self.assertEqual(self.client.get(f"/api/tasks/assignees/?part_of={self.main.id}").status_code, 400)

    def test_cancel_request_approved_cancels_the_parts(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        cr = TaskChangeRequest.objects.create(task=self.main, requested_by=self.rahul,
                                              changes={"cancel": True}, reason="client dropped it")
        self.as_(self.manager)
        res = self.client.post(f"/api/task-change-requests/{cr.id}/review/", {"decision": "approved"},
                               format="json")
        self.assertEqual(res.status_code, 200, res.data)
        self.assertIsNotNone(Task.objects.get(pk=p.pk).deleted_at)

    def test_approved_earlier_date_flags_the_late_part(self):
        self.part(self.give(self.rahul, self.main, self.amit, due_in_days=4))
        cr = TaskChangeRequest.objects.create(
            task=self.main, requested_by=self.rahul, reason="client wants it sooner",
            changes={"due_at": (self.now + timedelta(days=3)).isoformat()})
        self.as_(self.manager)
        self.client.post(f"/api/task-change-requests/{cr.id}/review/", {"decision": "approved"}, format="json")
        self.assertTrue(Notification.objects.filter(user=self.rahul, type="subtask_due_after_parent").exists())

    def test_new_notifications_open_the_task_and_sit_under_tasks(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        self.complete(self.amit, p)
        self.accept(self.rahul, p)
        n = Notification.objects.get(user=self.rahul, type="subtasks_all_done")
        self.assertEqual((n.category, n.link), ("tasks", f"/tasks/{self.main.id}"))

    def test_part_completion_waits_for_the_person_who_gave_it(self):
        p = self.part(self.give(self.rahul, self.main, self.amit))
        deeper = self.part(self.give(self.amit, p, self.vikram, effort=10))
        self.complete(self.vikram, deeper)
        self.assertEqual(TaskCompletion.objects.get(task=deeper).approver_id, self.amit.id)
        self.assertEqual(self.complete(self.amit, p).status_code, 400)   # deeper part not accepted yet
        self.accept(self.amit, deeper)
        self.assertEqual(self.complete(self.amit, p).status_code, 200)

    def test_effort_split_in_team_report_and_person_numbers(self):
        p = self.part(self.give(self.rahul, self.main, self.amit, effort=60))
        self.complete(self.amit, p)
        self.accept(self.rahul, p)
        self.complete(self.rahul, self.main)
        self.as_(self.admin)
        rows = {r["user"]: r for r in self.client.get("/api/tasks/team_report/?range=all").data["rows"]}
        self.assertEqual(rows[self.rahul.id]["time_earned_minutes"], 180)
        self.assertEqual(rows[self.amit.id]["time_earned_minutes"], 60)
        time = {r["user"]: r for r in self.client.get("/api/tasks/time_report/?range=all").data["rows"]}
        self.assertEqual((time[self.rahul.id]["time_earned_minutes"], time[self.amit.id]["time_earned_minutes"]),
                         (180, 60))

    def test_chains_tab_needs_the_company_wide_view(self):
        self.part(self.give(self.rahul, self.main, self.amit))
        for who in (self.rahul, self.manager, self.hr):
            self.as_(who)
            self.assertEqual(self.client.get("/api/tasks/chains/").status_code, 403, who.username)
            self.assertEqual(self.client.get(f"/api/tasks/{self.main.id}/chain/").status_code, 403, who.username)
        self.as_(self.admin)
        self.assertEqual(self.client.get("/api/tasks/chains/?status=bad").status_code, 400)
        self.assertEqual(len(self.client.get("/api/tasks/chains/?status=done").data["rows"]), 0)
        self.assertEqual(len(self.client.get("/api/tasks/chains/?status=all").data["rows"]), 1)


class ChangeRequestOldDueTests(ChainBase):
    """The ranking (Team performance) and the person Excel report show old -> new due date and the outcome."""

    def test_old_and_new_date_survive_the_approval(self):
        old = self.main.due_at
        new = self.now + timedelta(days=8)
        cr = TaskChangeRequest.objects.create(task=self.main, requested_by=self.rahul,
                                              changes={"due_at": new.isoformat()}, reason="supplier late")
        self.assertEqual(cr.old_due_at, old)
        self.as_(self.manager)
        self.client.post(f"/api/task-change-requests/{cr.id}/review/",
                         {"decision": "approved", "remarks": "ok"}, format="json")
        self.as_(self.admin)
        row = self.client.get("/api/tasks/change_request_leaders/?range=all").data["rows"][0]
        req = row["requests"][0]
        self.assertEqual((req["status"], req["reason"], req["remarks"]), (ChangeRequestStatus.APPROVED, "supplier late", "ok"))
        self.assertEqual(req["old_due_at"], old)
        self.assertEqual(req["new_due_at"].replace(microsecond=0), new.replace(microsecond=0))
        self.assertTrue(row["last_at"])

        res = self.client.get(f"/api/tasks/person_export/?user={self.rahul.id}&range=all")
        ws = load_workbook(BytesIO(res.content))["Change requests"]
        head, first = list(ws.iter_rows(values_only=True))[:2]
        self.assertIsNotNone(first[head.index("Old due")])
        self.assertIsNotNone(first[head.index("New due")])
        self.assertEqual(first[head.index("Status")], "Approved")

    def test_ranking_counts_and_permission(self):
        due = (self.now + timedelta(days=9)).isoformat()
        for status in (ChangeRequestStatus.APPROVED, ChangeRequestStatus.REJECTED, ChangeRequestStatus.PENDING):
            TaskChangeRequest.objects.create(task=self.main, requested_by=self.rahul, changes={"due_at": due},
                                             reason="r", status=status)
        TaskChangeRequest.objects.create(task=self.main, requested_by=self.rahul, changes={"title": "x"}, reason="r")
        for who in (self.rahul, self.manager):
            self.as_(who)
            self.assertEqual(self.client.get("/api/tasks/change_request_leaders/?range=all").status_code, 403)
        self.as_(self.admin)
        top = self.client.get("/api/tasks/change_request_leaders/?range=all").data["rows"][0]
        self.assertEqual((top["total"], top["approved"], top["rejected"], top["pending"], top["tasks"]),
                         (3, 1, 1, 1, 1))
