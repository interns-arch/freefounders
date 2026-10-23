"""Downloads everywhere (PLAN.md §19.31), the Tasks area: every list answers ?export=xlsx|csv with ALL the rows the
screen's filters match (not one page), with the list's own permission checks, and typed text never becomes a
live formula."""
import csv
import io
from datetime import timedelta

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Role, User
from crm.models import Lead, Task, TaskActivity, TaskChangeRequest, TaskCompletion, TaskTemplate
from intake.models import InboundMessage
from mistakes.models import SOP, Mistake, MistakeCategory
from notifications.models import Notification
from workspace.models import Group, Link, LinkCollection, Notice


def make(username, role, department="sales", **extra):
    return User.objects.create_user(username, f"{username}@x.com", "pass@12345", role=role,
                                    department=department, first_name=username.capitalize(), **extra)


def xlsx_rows(res):
    ws = load_workbook(io.BytesIO(res.content)).active
    return [[c.value for c in r] for r in ws.iter_rows()]


def csv_rows(res):
    return list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))


class ExportBase(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = make("boss", Role.ADMIN, "management")
        self.manager = make("meera", Role.SALES_MANAGER)
        self.rahul = make("rahul", Role.SALES_EXECUTIVE, reporting_manager=None)
        self.amit = make("amit", Role.SALES_EXECUTIVE)
        self.rahul.reporting_manager = self.manager
        self.rahul.save()
        self.amit.reporting_manager = self.manager
        self.amit.save()

    def as_(self, user):
        self.client.force_authenticate(user)
        return self.client

    def get(self, user, path, **params):
        return self.as_(user).get(path, params)

    def task(self, title, who, by=None, **kw):
        kw.setdefault("due_at", timezone.now() + timedelta(days=2))
        return Task.objects.create(title=title, assigned_to=who, created_by=by or self.manager, effort_minutes=30, **kw)


class TaskListExportTests(ExportBase):
    def setUp(self):
        super().setUp()
        for i in range(30):   # more than one page (25)
            self.task(f"Call client {i}", self.rahul)
        self.task("=HYPERLINK(\"http://evil\")", self.rahul)
        self.task("Done thing", self.rahul, status="done", completed_at=timezone.now())
        self.task("Amit's task", self.amit)

    def test_all_matching_rows_not_one_page_and_filters_applied(self):
        res = self.get(self.rahul, "/api/tasks/", scope="my", status="open,in_progress", export="xlsx")
        self.assertEqual(res.status_code, 200, res.content[:300])
        self.assertIn("spreadsheetml", res["Content-Type"])
        self.assertIn("attachment", res["Content-Disposition"])
        rows = xlsx_rows(res)
        self.assertEqual(rows[0][:4], ["Code", "Title", "Assigned to", "Assigned by"])
        self.assertEqual(len(rows) - 1, 31)            # 30 + the formula one; not the done one, not Amit's
        titles = [r[1] for r in rows[1:]]
        self.assertNotIn("Done thing", titles)
        self.assertNotIn("Amit's task", titles)
        # the page itself is still paginated
        self.assertEqual(len(self.get(self.rahul, "/api/tasks/", scope="my").data["results"]), 25)

    def test_formula_neutralised_and_csv(self):
        res = self.get(self.rahul, "/api/tasks/", scope="my", search="HYPERLINK", export="csv")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res["Content-Type"].startswith("text/csv"))
        rows = csv_rows(res)
        self.assertEqual(rows[0][0], "Code")
        self.assertEqual(len(rows), 2)
        self.assertTrue(rows[1][1].startswith("'="), rows[1][1])
        x = xlsx_rows(self.get(self.rahul, "/api/tasks/", scope="my", search="HYPERLINK", export="xlsx"))
        self.assertTrue(x[1][1].startswith("'="))

    def test_permission_is_the_lists(self):
        # Amit sees only his own task in "All tasks", never Rahul's
        rows = xlsx_rows(self.get(self.amit, "/api/tasks/", export="xlsx"))
        self.assertEqual([r[1] for r in rows[1:]], ["Amit's task"])
        # the Deleted bin is refused without its permission, file or not
        self.assertEqual(self.get(self.amit, "/api/tasks/", scope="deleted", export="xlsx").status_code, 403)
        t = Task.objects.get(title="Amit's task")
        t.deleted_at = timezone.now()
        t.save()
        rows = xlsx_rows(self.get(self.admin, "/api/tasks/", scope="deleted", export="xlsx"))
        self.assertEqual([r[1] for r in rows[1:]], ["Amit's task"])

    def test_unknown_format_is_the_plain_list(self):
        res = self.get(self.rahul, "/api/tasks/", scope="my", export="pdf")
        self.assertEqual(res.status_code, 200)
        self.assertIn("results", res.data)


class TaskInboxExportTests(ExportBase):
    def test_reviews_filtered_by_person_and_only_mine(self):
        t1, t2 = self.task("R1", self.rahul), self.task("A1", self.amit)
        TaskCompletion.objects.create(task=t1, submitted_by=self.rahul, approver=self.manager, note="=1+1")
        TaskCompletion.objects.create(task=t2, submitted_by=self.amit, approver=self.manager)
        TaskCompletion.objects.create(task=t2, submitted_by=self.amit, approver=self.admin)
        rows = xlsx_rows(self.get(self.manager, "/api/task-completions/", scope="inbox", export="xlsx"))
        self.assertEqual(len(rows) - 1, 2)                       # not the one waiting for the admin
        rows = xlsx_rows(self.get(self.manager, "/api/task-completions/", scope="inbox",
                                  submitted_by=self.rahul.id, export="xlsx"))
        self.assertEqual([r[1] for r in rows[1:]], ["R1"])
        self.assertEqual(rows[1][3], "'=1+1")
        self.assertEqual(self.get(self.rahul, "/api/task-completions/", scope="all", export="xlsx").status_code, 403)

    def test_change_requests(self):
        t = self.task("Change me", self.rahul)
        TaskChangeRequest.objects.create(task=t, requested_by=self.rahul, changes={"title": "x"}, reason="-oops")
        rows = xlsx_rows(self.get(self.manager, "/api/task-change-requests/", scope="inbox", export="xlsx"))
        self.assertEqual(rows[0][:3], ["Code", "Task", "Assignee"])
        self.assertEqual([r[1] for r in rows[1:]], ["Change me"])
        self.assertEqual(rows[1][6], "'-oops")
        self.assertEqual(len(xlsx_rows(self.get(self.amit, "/api/task-change-requests/", scope="inbox", export="xlsx"))), 1)

    def test_activities_and_templates(self):
        mine, other = self.task("Mine", self.rahul), self.task("Theirs", self.amit)
        TaskActivity.objects.create(task=mine, actor=self.rahul, text="@started", kind="comment")
        TaskActivity.objects.create(task=other, actor=self.amit, text="hidden", kind="comment")
        rows = xlsx_rows(self.get(self.rahul, "/api/task-activities/", export="xlsx"))
        self.assertEqual([r[6] for r in rows[1:]], ["'@started"])
        TaskTemplate.objects.create(name="Daily", title="Do it", created_by=self.admin)
        rows = csv_rows(self.get(self.rahul, "/api/task-templates/", export="csv"))
        self.assertEqual(rows[0][0], "Name")
        self.assertEqual(rows[1][0], "Daily")


class WorkspaceExportTests(ExportBase):
    def test_groups_notices_links(self):
        g = Group.objects.create(name="Mine", owner=self.manager)
        g.members.add(self.rahul)
        Group.objects.create(name="Other", owner=self.manager)
        rows = xlsx_rows(self.get(self.rahul, "/api/groups/", export="xlsx"))
        self.assertEqual([r[0] for r in rows[1:]], ["Mine"])
        self.assertIn("Rahul", rows[1][5])

        Notice.objects.create(title="=cmd", content="hi", author=self.admin, status="published",
                              publish_at=timezone.now() - timedelta(hours=1))
        res = self.get(self.rahul, "/api/notices/", export="xlsx")
        self.assertEqual(res.status_code, 200, res.content[:300])
        rows = xlsx_rows(res)
        self.assertEqual(rows[0][0], "Title")
        self.assertEqual([r[0] for r in rows[1:]], ["'=cmd"])
        self.assertEqual(self.get(self.rahul, "/api/notices/", manage="true", export="xlsx").status_code, 403)

        c = LinkCollection.objects.create(name="Tools")
        Link.objects.create(collection=c, title="Drive", url="https://drive.example.com", added_by=self.admin)
        Link.objects.create(collection=c, title="Wiki", url="https://wiki.example.com", added_by=self.admin)
        rows = xlsx_rows(self.get(self.rahul, "/api/links/", search="wiki", export="xlsx"))
        self.assertEqual([r[0] for r in rows[1:]], ["Wiki"])
        self.assertEqual(rows[1][1], "https://wiki.example.com")


class OtherAreaExportTests(ExportBase):
    def test_leads_are_only_the_visible_ones(self):
        Lead.objects.create(customer_name="Ravi", department="sales", assigned_to=self.rahul, phone="+919800000000")
        Lead.objects.create(customer_name="Sunita", department="sales", assigned_to=self.amit)
        rows = xlsx_rows(self.get(self.rahul, "/api/leads/", export="xlsx"))
        self.assertEqual([r[0] for r in rows[1:]], ["Ravi"])
        self.assertEqual(rows[1][2], "'+919800000000")     # a leading + is neutralised too
        rows = xlsx_rows(self.get(self.manager, "/api/leads/", search="sun", export="xlsx"))
        self.assertEqual([r[0] for r in rows[1:]], ["Sunita"])

    def test_ai_inbox_needs_its_permission(self):
        InboundMessage.objects.create(channel="wa", external_id="1", sender="98", body="Need 3 tyres")
        InboundMessage.objects.create(channel="gmail", external_id="2", sender="a@b.c", body="Hello")
        self.assertEqual(self.get(self.rahul, "/api/intake/", export="xlsx").status_code, 403)
        rows = xlsx_rows(self.get(self.admin, "/api/intake/", channel="wa", export="xlsx"))
        self.assertEqual([r[5] for r in rows[1:]], ["Need 3 tyres"])

    def test_mistakes_own_only_and_sops(self):
        MistakeCategory.objects.get_or_create(name="Billing", defaults={"active": True})
        Mistake.objects.create(employee=self.rahul, reported_by=self.manager, category="Billing", severity="high",
                               description="=SUM(A1)", department="sales")
        Mistake.objects.create(employee=self.amit, reported_by=self.manager, category="Billing", severity="low",
                               description="Amit's", department="sales")
        rows = xlsx_rows(self.get(self.rahul, "/api/mistakes/", export="xlsx"))
        self.assertEqual([r[6] for r in rows[1:]], ["'=SUM(A1)"])
        rows = xlsx_rows(self.get(self.admin, "/api/mistakes/", severity="low", export="xlsx"))
        self.assertEqual([r[6] for r in rows[1:]], ["Amit's"])

        SOP.objects.create(title="Billing", department="accounts", version="v1", steps="1. Check", active=True)
        SOP.objects.create(title="Old", department="accounts", version="v0", steps="x", active=False)
        rows = xlsx_rows(self.get(self.rahul, "/api/sops/", export="xlsx"))
        self.assertEqual([r[0] for r in rows[1:]], ["Billing"])
        self.assertEqual(len(xlsx_rows(self.get(self.rahul, "/api/sops/", all="true", export="xlsx"))) - 1, 2)

    def test_notifications_are_mine_only(self):
        for i in range(60):
            Notification.objects.create(user=self.rahul, type="x", title=f"n{i}")
        Notification.objects.create(user=self.amit, type="x", title="not yours")
        rows = xlsx_rows(self.get(self.rahul, "/api/notifications/", export="xlsx"))
        self.assertEqual(len(rows) - 1, 60)
        self.assertNotIn("not yours", [r[2] for r in rows[1:]])
