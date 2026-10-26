"""Per-person Excel performance report (original app 5fde3aa; PLAN.md §19.41): Admins and the Owner only
(company-wide task view), and the numbers in the file match the screens (same score helper)."""
from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import Role
from .models import Task, TaskChangeRequest, TaskStatus
from .tests_task_engine import Base, make

URL = "/api/tasks/person_export/"


class PersonExportBase(Base):
    def setUp(self):
        super().setUp()
        self.now = timezone.now()
        self.super = make("chief", Role.SUPER_ADMIN, "management")

    def task(self, who, giver, *, done=False, late=False, overdue=False, effort=60):
        due = self.now - timedelta(days=2) if (done or overdue) else self.now + timedelta(days=2)
        t = Task.objects.create(title="T", assigned_to=who, created_by=giver,
                                effort_minutes=effort, due_at=due,
                                status=TaskStatus.DONE if done else TaskStatus.OPEN)
        if done:
            Task.objects.filter(pk=t.pk).update(
                completed_at=due + timedelta(hours=5 if late else -5))
        return t

    def download(self, user, **params):
        q = "&".join(f"{k}={v}" for k, v in {"user": user.pk, "range": "all", **params}.items())
        return self.client.get(f"{URL}?{q}")

    def sheet(self, res, name):
        wb = load_workbook(BytesIO(res.content))
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        return ws, rows[0], rows[1:]


class PersonExportTests(PersonExportBase):
    def test_only_admin_and_super_admin(self):
        for who in (self.rahul, self.manager, self.hr):
            self.as_(who)
            self.assertEqual(self.download(self.amit).status_code, 403, who.username)
        for who in (self.admin, self.super):
            self.as_(who)
            res = self.download(self.amit)
            self.assertEqual(res.status_code, 200, who.username)
            self.assertIn("spreadsheetml", res["Content-Type"])
            self.assertIn(".xlsx", res["Content-Disposition"])

    def test_unknown_user(self):
        self.as_(self.admin)
        self.assertEqual(self.client.get(f"{URL}?user=999999").status_code, 400)
        self.assertEqual(self.client.get(f"{URL}?user=abc").status_code, 400)

    def test_all_sheets_with_filters_and_frozen_header(self):
        self.task(self.rahul, self.manager, done=True)
        self.as_(self.admin)
        wb = load_workbook(BytesIO(self.download(self.rahul).content))
        self.assertEqual(wb.sheetnames, [
            "Summary", "Month by month", "Tasks received", "Given by whom",
            "Tasks delegated", "Delegated to whom", "Change requests",
            "Completion reviews"])
        ws = wb["Tasks received"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.auto_filter.ref)

    def test_score_matches_the_employees_report(self):
        self.task(self.rahul, self.manager, done=True)
        self.task(self.rahul, self.manager, done=True, late=True)
        self.task(self.rahul, self.admin, overdue=True)
        self.as_(self.admin)
        screen = self.client.get(
            f"/api/tasks/employees_report/?range=all&user={self.rahul.pk}").data["rows"][0]
        ws, _, _ = self.sheet(self.download(self.rahul), "Summary")
        summary = {r[0]: r[1] for r in ws.iter_rows(values_only=True)}
        self.assertEqual(summary["SCORE"], screen["score"])
        self.assertEqual(summary["Completed"], 2)
        self.assertEqual(summary["Overdue"], 1)

    def test_tasks_received_show_who_gave_them_and_the_result(self):
        self.task(self.rahul, self.manager, done=True)
        self.task(self.rahul, self.manager, done=True, late=True)
        self.task(self.rahul, self.admin, overdue=True)
        self.task(self.rahul, self.rahul)                       # self-assigned
        self.as_(self.admin)
        res = self.download(self.rahul)
        _, head, rows = self.sheet(res, "Tasks received")
        giver, result, scored = head.index("Given by"), head.index("Result"), head.index("Scored")
        self.assertEqual(sorted(r[result] for r in rows),
                         ["Delayed", "On time", "Overdue", "Pending"])
        self.assertEqual(sum(1 for r in rows if r[scored] == "No"), 1)

        _, head, rows = self.sheet(res, "Given by whom")
        by = {r[0]: r for r in rows}
        meera = by[self.manager.get_full_name() or self.manager.username]
        self.assertEqual(meera[head.index("Tasks")], 2)
        self.assertEqual(meera[head.index("On time")], 1)
        self.assertEqual(meera[head.index("Delayed")], 1)

    def test_tasks_delegated_and_to_whom(self):
        self.task(self.amit, self.rahul, done=True)
        self.task(self.vikram, self.rahul, overdue=True)
        self.as_(self.admin)
        res = self.download(self.rahul)
        _, head, rows = self.sheet(res, "Tasks delegated")
        self.assertEqual(len(rows), 2)
        _, head, rows = self.sheet(res, "Delegated to whom")
        self.assertEqual(len(rows), 2)

    def test_change_requests_sheet_and_range(self):
        t = self.task(self.rahul, self.manager)
        TaskChangeRequest.objects.create(task=t, requested_by=self.rahul,
                                         changes={"due_at": self.now.isoformat()},
                                         reason="supplier late")
        old = TaskChangeRequest.objects.create(task=t, requested_by=self.rahul,
                                               changes={"cancel": True}, reason="x")
        TaskChangeRequest.objects.filter(pk=old.pk).update(
            created_at=self.now - timedelta(days=400))
        self.as_(self.admin)
        _, head, rows = self.sheet(self.download(self.rahul), "Change requests")
        self.assertEqual(sorted(r[head.index("Type")] for r in rows),
                         ["Cancel", "Due date / time"])
        _, head, rows = self.sheet(self.download(self.rahul, range="this_month"),
                                   "Change requests")
        self.assertEqual([r[head.index("Type")] for r in rows], ["Due date / time"])

    def test_person_with_nothing_still_downloads(self):
        self.as_(self.super)
        res = self.download(self.vikram, range="this_week")
        self.assertEqual(res.status_code, 200)
        _, _, rows = self.sheet(res, "Tasks received")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], "Nothing in this range.")


class PersonExportKaampilotTests(PersonExportBase):
    """Kaampilot checks on top of the original's (PLAN.md §19.41): every sheet filterable, the Summary equal
    to /tasks/employees_report/?user= (the API behind the person's profile on screen) field by field."""

    def test_every_sheet_has_filter_arrows(self):
        self.task(self.rahul, self.manager, done=True)
        self.as_(self.admin)
        wb = load_workbook(BytesIO(self.download(self.rahul).content))
        self.assertEqual(len(wb.sheetnames), 8)
        for ws in wb.worksheets:
            self.assertTrue(ws.auto_filter.ref, ws.title)

    def test_summary_equals_the_screen_api(self):
        self.task(self.rahul, self.manager, done=True, effort=90)
        self.task(self.rahul, self.manager, done=True, late=True, effort=30)
        self.task(self.rahul, self.admin, overdue=True, effort=45)
        self.task(self.rahul, self.rahul, effort=20)
        self.as_(self.admin)
        for rng in ("all", "this_month"):
            screen = self.client.get(
                f"/api/tasks/employees_report/?range={rng}&user={self.rahul.pk}").data["rows"][0]
            ws, _, _ = self.sheet(self.download(self.rahul, range=rng), "Summary")
            summary = {r[0]: r[1] for r in ws.iter_rows(values_only=True)}
            pairs = {"SCORE": "score", "Total in range": "total", "Completed": "completed",
                     "  on time": "in_time", "  delayed": "delayed", "Pending": "pending",
                     "In progress": "in_progress", "Overdue": "overdue",
                     "Self-assigned / correction (not scored)": "self_assigned",
                     "On-time rate % (scored tasks)": "on_time_rate",
                     "Effort ratio % (earned / assigned)": "effort_ratio"}
            for label, key in pairs.items():
                want = screen[key] if screen[key] is not None else "-"
                self.assertEqual(summary[label], want, f"{rng}: {label}")
            self.assertEqual(summary["Time earned (h)"], round(screen["time_earned_minutes"] / 60, 2))
            self.assertEqual(summary["Time assigned (h)"], round(screen["time_assigned_minutes"] / 60, 2))

    def test_owner_without_admin_role_and_formula_safe_text(self):
        t = self.task(self.rahul, self.manager)
        Task.objects.filter(pk=t.pk).update(title="=HYPERLINK(1)")
        self.as_(self.super)
        _, head, rows = self.sheet(self.download(self.rahul), "Tasks received")
        self.assertEqual(rows[0][head.index("Title")], "'=HYPERLINK(1)")
