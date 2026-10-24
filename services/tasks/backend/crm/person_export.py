"""One person's performance as an Excel workbook, for Admins and the Owner (original app 5fde3aa + the chain columns
of 1dabd35; PLAN.md §19.41).

Every number comes from the same helpers the screens use (the score is TaskViewSet._person_rows, also behind
/tasks/employees_report/?user= that the person's profile shows), so the file can never disagree with the page.
Each sheet is a plain table with filters switched on and the header frozen, so it can be sliced in Excel any way
the reviewer likes: by giver, status, result, month, priority, category, chain ...
Text typed by people is written formula-safe (config.exports.safe_cell, PLAN.md §19.31).
"""
from collections import OrderedDict
from datetime import date, timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.exports import safe_cell

from .models import Task, TaskChangeRequest, TaskCompletion, TaskStatus

SHEETS = ["Summary", "Month by month", "Tasks received", "Given by whom", "Tasks delegated",
          "Delegated to whom", "Change requests", "Completion reviews"]
HEAD_FILL = PatternFill("solid", fgColor="4F46E5")      # the brand indigo
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT_COLOUR = "3730A3"
RESULT_FILL = {
    "On time": PatternFill("solid", fgColor="E3F3ED"),
    "Delayed": PatternFill("solid", fgColor="FDF1E3"),
    "Overdue": PatternFill("solid", fgColor="FDECEA"),
}
DT_FMT = "dd mmm yyyy hh:mm"
EMPTY = "Nothing in this range."


def _local(dt):
    """Excel has no time zones: write the company's wall-clock time."""
    return timezone.localtime(dt).replace(tzinfo=None) if dt else None


def _name(u):
    return (u.get_full_name() or u.username) if u else ""


def _role(u):
    try:
        return u.get_role_display() if u else ""
    except Exception:      # noqa: BLE001 -- a label is never worth a failed download
        return ""


def _hours(minutes):
    return round(minutes / 60, 2) if minutes else 0


def task_result(t, now):
    """On time / Delayed / Overdue / In progress / Pending, plus days late."""
    if t.status == TaskStatus.DONE:
        if not t.due_at or not t.completed_at:
            return "Done (no due date)", None
        if t.completed_at > t.due_at:
            return "Delayed", round((t.completed_at - t.due_at).total_seconds() / 86400, 1)
        return "On time", None
    if t.due_at and t.due_at < now:
        return "Overdue", round((now - t.due_at).total_seconds() / 86400, 1)
    return ("In progress" if t.status == TaskStatus.IN_PROGRESS else "Pending"), None


def _in_range(t, start, end):
    """The reports' anchor: the due date, else the created date (same as _task_buckets)."""
    if not start:
        return True
    anchor = timezone.localtime(t.due_at or t.created_at).date()
    return start <= anchor <= end


def _month_bounds(first):
    nxt = (first + timedelta(days=32)).replace(day=1)
    return first, nxt - timedelta(days=1)


def _table(ws, header, rows, widths=None, result_col=None):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([safe_cell(v) for v in r])
    for col in ws.iter_cols(min_row=2):
        for c in col:
            if hasattr(c.value, "year") and hasattr(c.value, "hour"):
                c.number_format = DT_FMT
            elif isinstance(c.value, date):
                c.number_format = "dd mmm yyyy"
    if result_col is not None:
        for (c,) in ws.iter_rows(min_row=2, min_col=result_col, max_col=result_col):
            if c.value in RESULT_FILL:
                c.fill = RESULT_FILL[c.value]
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h) or min(max(len(str(h)) + 2, 12), 40)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if not rows:
        ws.append([EMPTY])
    # filter arrows on the header row of every table sheet, also when it is empty
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(ws.max_row, 2)}"


TASK_HEADER_TAIL = ["Status", "Result", "Days late", "Priority", "Recurrence",
                    "Category", "Group", "Created", "Due", "Completed",
                    "Effort (h)", "Actual (h)", "Scored", "Approval", "Month",
                    "Sub-task of", "Main task", "Given to sub-tasks (h)"]
TASK_WIDTHS = {"Title": 42, "Created": 18, "Due": 18, "Completed": 18}


def _task_rows(tasks, corrective, now, other):
    """`other` picks the column for the other side: who gave it / who got it."""
    rows = []
    for t in tasks:
        result, days_late = task_result(t, now)
        own = t.created_by_id == t.assigned_to_id or t.id in corrective
        who = other(t)
        review = t._review.get_status_display() if getattr(t, "_review", None) else ""
        rows.append([
            t.code, t.title, _name(who), _role(who),
            t.get_status_display(), result, days_late,
            t.get_priority_display(), t.get_frequency_display(), t.category,
            t.group.name if t.group_id else "",
            _local(t.created_at), _local(t.due_at), _local(t.completed_at),
            _hours(t.effort_minutes), _hours(t.actual_minutes),
            "No" if own else "Yes", review,
            timezone.localtime(t.due_at or t.created_at).strftime("%Y-%m"),
            getattr(t, "_part_of", ""), getattr(t, "_main", ""),
            _hours(getattr(t, "_given", 0)),
        ])
    return rows


def _by_person(tasks, now, side):
    """One row per giver (or receiver): how that relationship is going."""
    agg = OrderedDict()
    for t in tasks:
        who = side(t)
        key = who.pk if who else 0
        a = agg.setdefault(key, {"name": _name(who) or "(nobody)", "role": _role(who),
                                 "total": 0, "done": 0, "on_time": 0, "delayed": 0,
                                 "overdue": 0, "open": 0, "effort": 0, "earned": 0})
        result, _ = task_result(t, now)
        a["total"] += 1
        a["effort"] += t.effort_minutes or 0
        if t.status == TaskStatus.DONE:
            a["done"] += 1
            a["earned"] += t.effort_minutes or 0
            if result == "On time":
                a["on_time"] += 1
            elif result == "Delayed":
                a["delayed"] += 1
        else:
            a["open"] += 1
            if result == "Overdue":
                a["overdue"] += 1
    rows = []
    for a in sorted(agg.values(), key=lambda a: -a["total"]):
        judged = a["done"] + a["overdue"]
        rows.append([a["name"], a["role"], a["total"], a["done"], a["on_time"],
                     a["delayed"], a["overdue"], a["open"],
                     round(100 * a["on_time"] / judged, 1) if judged else None,
                     _hours(a["effort"]), _hours(a["earned"])])
    return rows


BY_PERSON_TAIL = ["Tasks", "Completed", "On time", "Delayed", "Overdue now",
                  "Open", "On-time %", "Effort assigned (h)", "Effort earned (h)"]


def build_person_workbook(view, person, start, end, range_label, viewer):
    """Returns (bytes, filename)."""
    from .chain import own_effort_map, root_of
    from .scoring import corrective_task_ids
    now = timezone.now()
    today = timezone.localtime(now).date()
    name = _name(person)

    # ---- the score, exactly as the screens compute it -------------------
    per, corrective = view._task_buckets([person], start, end)
    s = view._person_rows(per, corrective, {person.pk: name}, start, end, now, today, always_emit=True)[0]

    base = (Task.objects.filter(deleted_at__isnull=True)
            .select_related("created_by", "assigned_to", "group")
            .order_by("due_at", "id"))
    received = [t for t in base.filter(assigned_to=person) if _in_range(t, start, end)]
    delegated = [t for t in base.filter(created_by=person).exclude(assigned_to=person)
                 if _in_range(t, start, end)]
    corrective_all = corrective_task_ids([t.id for t in received + delegated])

    # Task chain: effort handed to sub-tasks is earned by those people, so "Effort" here is what the holder kept
    # (same rule as the score); each row says which task it is a part of and the chain's main task.
    own_map = own_effort_map(t.id for t in received + delegated)
    roots = {}
    for t in received + delegated:
        if t.id in own_map:
            t._given = (t.effort_minutes or 0) - own_map[t.id]
            t.effort_minutes = own_map[t.id]          # in memory only, never saved
        if t.parent_id:
            if t.parent_id not in roots:
                parent = Task.objects.only("id", "parent_id").get(pk=t.parent_id)
                roots[t.parent_id] = (parent.code, root_of(parent).code)
            t._part_of, t._main = roots[t.parent_id]

    # latest approval decision per task, for the "Approval" column
    latest = {}
    for c in (TaskCompletion.objects.filter(task__in=[t.id for t in received + delegated])
              .order_by("created_at")):
        latest[c.task_id] = c
    for t in received + delegated:
        t._review = latest.get(t.id)

    crs = TaskChangeRequest.objects.filter(requested_by=person) \
        .select_related("task", "reviewed_by").order_by("-created_at")
    subs = TaskCompletion.objects.filter(submitted_by=person) \
        .select_related("task", "approver", "reviewed_by").order_by("-created_at")
    if start:
        crs = crs.filter(created_at__date__range=(start, end))
        subs = subs.filter(created_at__date__range=(start, end))
    crs, subs = list(crs), list(subs)

    wb = Workbook()

    # ---- 1. Summary ------------------------------------------------------
    ws = wb.active
    ws.title = SHEETS[0]
    period = f"{start:%d %b %Y} - {end:%d %b %Y}" if start else "All time"
    pct = lambda v: v if v is not None else "-"   # noqa: E731  (a number, the % is in the label)
    deadline_asks = sum(1 for c in crs if "due_at" in (c.changes or {}))
    department = ""
    try:
        department = person.get_department_display()
    except Exception:      # noqa: BLE001
        department = getattr(person, "department", "") or ""
    lines = [
        ("Performance report", name),
        ("Role", _role(person)),
        ("Department", department),
        ("Reporting manager", _name(getattr(person, "reporting_manager", None))),
        ("Period", f"{range_label} ({period})" if start else "All time"),
        ("Generated", f"{timezone.localtime(now):%d %b %Y %I:%M %p} by {_name(viewer)}"),
        ("", ""),
        ("SCORE", s["score"] if s["score"] is not None else "-"),
        ("Score before mistake penalty", s.get("task_score") if s.get("task_score") is not None else "-"),
        ("Mistake penalty", -(s.get("mistake_penalty") or 0)),
        ("On-time rate % (scored tasks)", pct(s["on_time_rate"])),
        ("Effort ratio % (earned / assigned)", pct(s["effort_ratio"])),
        ("Review", s.get("review") or ""),
        ("", ""),
        ("TASKS RECEIVED", ""),
        ("Total in range", s["total"]),
        ("Completed", s["completed"]),
        ("  on time", s["in_time"]),
        ("  delayed", s["delayed"]),
        ("Pending", s["pending"]),
        ("In progress", s["in_progress"]),
        ("Overdue", s["overdue"]),
        ("Self-assigned / correction (not scored)", s["self_assigned"]),
        ("Given by others (scored)", s["total"] - s["self_assigned"]),
        ("Time assigned (h)", _hours(s["time_assigned_minutes"])),
        ("Time earned (h)", _hours(s["time_earned_minutes"])),
        ("Time actually spent (h)", _hours(s["time_spent_minutes"])),
        ("Multitask days", s.get("multitask_days", 0)),
        ("Sub-tasks received (parts of someone's task)", sum(1 for t in received if t.parent_id)),
        ("", ""),
        ("TASKS DELEGATED (given to others)", ""),
        ("Tasks given to others", len(delegated)),
        ("  of those completed", sum(1 for t in delegated if t.status == TaskStatus.DONE)),
        ("  of those overdue now", sum(1 for t in delegated if task_result(t, now)[0] == "Overdue")),
        ("  of those sub-tasks (parts given)", sum(1 for t in delegated if t.parent_id)),
        ("People given work", len({t.assigned_to_id for t in delegated})),
        ("", ""),
        ("QUALITY & REQUESTS", ""),
        ("Completions submitted", len(subs)),
        ("  sent back for redo", sum(1 for c in subs if c.status == "rejected")),
        ("Mistakes logged", s.get("mistakes", 0)),
        ("  repeat mistakes", s.get("repeat_mistakes", 0)),
        ("Change requests raised", len(crs)),
        ("  to move a due date / time", deadline_asks),
        ("  to cancel a task", sum(1 for c in crs if (c.changes or {}).get("cancel"))),
        ("", ""),
        ("How the score works", view._score_formula()),
        ("How to use this file", "Every sheet has filters on the header row - click the "
                                 "arrows to slice by giver, status, result, month, "
                                 "priority, category, group or chain."),
    ]
    for k, v in lines:
        ws.append([k, safe_cell(v)])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 70
    for (a, b) in ws.iter_rows(min_row=1, max_col=2):
        b.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        if a.value and (a.value.isupper() or a.row == 1):
            a.font = Font(bold=True, color=TITLE_FONT_COLOUR, size=12 if a.row == 1 else 11)
    ws["B1"].font = Font(bold=True, size=14)
    ws.cell(row=[k for k, _ in lines].index("SCORE") + 1, column=2).font = Font(bold=True, size=14)
    ws.auto_filter.ref = f"A1:B{ws.max_row}"

    # ---- 2. Month by month -----------------------------------------------
    ws = wb.create_sheet(SHEETS[1])
    if start:
        first, last = start.replace(day=1), end.replace(day=1)
    else:
        # all time: from the month of this person's first task (max 36 months)
        last = today.replace(day=1)
        anchors = [timezone.localtime(t.due_at or t.created_at).date() for t in received + delegated]
        first = min([last, *anchors]).replace(day=1)
        if (last.year - first.year) * 12 + last.month - first.month >= 36:
            first = (date(last.year - 3, last.month, 1) + timedelta(days=32)).replace(day=1)
    months, m = [], first
    while m <= last and len(months) < 36:
        months.append(m)
        m = (m + timedelta(days=32)).replace(day=1)
    mrows = []
    for m in months:
        ms, me = _month_bounds(m)
        p, cor = view._task_buckets([person], ms, me)
        r = view._person_rows(p, cor, {person.pk: name}, ms, me, now, today, always_emit=True)[0]
        mrows.append([m.strftime("%b %Y"), r["score"], r["on_time_rate"],
                      r["total"], r["completed"], r["in_time"], r["delayed"],
                      r["overdue"], _hours(r["time_assigned_minutes"]),
                      _hours(r["time_earned_minutes"]), r.get("mistakes", 0)])
    _table(ws, ["Month", "Score", "On-time %", "Tasks", "Completed", "On time",
                "Delayed", "Overdue", "Assigned (h)", "Earned (h)", "Mistakes"], mrows)

    # ---- 3. Tasks received -----------------------------------------------
    ws = wb.create_sheet(SHEETS[2])
    _table(ws, ["Code", "Title", "Given by", "Giver role", *TASK_HEADER_TAIL],
           _task_rows(received, corrective_all, now, lambda t: t.created_by),
           widths=TASK_WIDTHS, result_col=6)

    # ---- 4. Who gave the work --------------------------------------------
    ws = wb.create_sheet(SHEETS[3])
    _table(ws, ["Given by", "Role", *BY_PERSON_TAIL],
           _by_person(received, now, lambda t: t.created_by), widths={"Given by": 26})

    # ---- 5. Tasks delegated ----------------------------------------------
    ws = wb.create_sheet(SHEETS[4])
    _table(ws, ["Code", "Title", "Given to", "Receiver role", *TASK_HEADER_TAIL],
           _task_rows(delegated, corrective_all, now, lambda t: t.assigned_to),
           widths=TASK_WIDTHS, result_col=6)

    # ---- 6. Who got the work ---------------------------------------------
    ws = wb.create_sheet(SHEETS[5])
    _table(ws, ["Given to", "Role", *BY_PERSON_TAIL],
           _by_person(delegated, now, lambda t: t.assigned_to), widths={"Given to": 26})

    # ---- 7. Change requests ----------------------------------------------
    ws = wb.create_sheet(SHEETS[6])
    crows = []
    for c in crs:
        ch = c.changes or {}
        kind = ("Cancel" if ch.get("cancel") else
                "Due date / time" if "due_at" in ch else "Other edit")
        old_due, new_due = c.due_change()
        crows.append([_local(c.created_at), c.task.code, c.task.title, kind,
                      _local(old_due), _local(new_due),
                      " | ".join(c.describe()), c.reason, c.get_status_display(),
                      _name(c.reviewed_by), _local(c.reviewed_at), c.remarks,
                      timezone.localtime(c.created_at).strftime("%Y-%m")])
    _table(ws, ["Asked on", "Code", "Task", "Type", "Old due", "New due", "Changes", "Reason", "Status",
                "Reviewed by", "Reviewed on", "Remarks", "Month"], crows,
           widths={"Asked on": 18, "Task": 36, "Old due": 18, "New due": 18, "Changes": 48, "Reason": 40,
                   "Reviewed on": 18, "Remarks": 30})

    # ---- 8. Completion reviews -------------------------------------------
    ws = wb.create_sheet(SHEETS[7])
    srows = [[_local(c.created_at), c.task.code, c.task.title, c.note,
              _name(c.approver), c.get_status_display(), _name(c.reviewed_by),
              _local(c.reviewed_at), c.remarks,
              timezone.localtime(c.created_at).strftime("%Y-%m")] for c in subs]
    _table(ws, ["Submitted", "Code", "Task", "Note", "Approver", "Status",
                "Reviewed by", "Reviewed on", "Remarks", "Month"], srows,
           widths={"Submitted": 18, "Task": 36, "Note": 40, "Reviewed on": 18, "Remarks": 30})

    buf = BytesIO()
    wb.save(buf)
    slug = lambda v: "-".join("".join(ch if ch.isalnum() else " " for ch in v.lower()).split())  # noqa: E731
    return buf.getvalue(), f"tasks-performance-{slug(name) or 'person'}-{slug(range_label)}-{today:%Y-%m-%d}.xlsx"
