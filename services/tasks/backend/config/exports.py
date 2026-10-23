"""Spreadsheet exports never carry live formulas.

Text people typed (a leave reason, a form answer, a name) that starts with = + - @ (or a tab / carriage return) is
run as a formula by Excel, LibreOffice and Google Sheets when the CSV/XLSX is opened: "CSV injection". Such a cell
is written with a leading apostrophe, which spreadsheets show as plain text. Numbers stay numbers.

Downloads everywhere (PLAN.md §19.31): every list endpoint that mixes in `ExportMixin` answers
`?export=xlsx|csv` (plus the screen's usual filters, search, sort and tab) with a file of ALL the matching rows
(up to EXPORT_MAX), built by the same queryset, permission checks and serializer as the page the screen shows.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal

FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
EXPORT_FORMATS = ("xlsx", "csv")
EXPORT_MAX = 50_000

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def safe_cell(value):
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def safe_row(row):
    return [safe_cell(v) for v in row]


# ----------------------------------------------------------------- writers

def to_csv(header, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(safe_row(header))
    for r in rows:
        w.writerow(safe_row([_csv_value(v) for v in r]))
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM: Excel opens it as UTF-8


def _csv_value(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return v


def to_xlsx(title, header, rows, extra_sheet=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", str(title))[:31] or "Report"
    ws.append(safe_row(header))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(safe_row([_xlsx_value(v) for v in r]))
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(4, width + 2))
    # dates and times show as dates in Excel
    for col in ws.iter_cols(min_row=2):
        for cell in col:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd-mm-yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd-mm-yyyy"
    if extra_sheet:
        name, lines = extra_sheet
        notes = wb.create_sheet(name)
        for line in lines:
            notes.append([line])
        notes.column_dimensions["A"].width = 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_value(v):
    if isinstance(v, datetime) and v.tzinfo is not None:
        from django.utils import timezone
        return timezone.localtime(v).replace(tzinfo=None)
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (list, dict, set, tuple)):
        return str(v)
    return v


def file_response(content: bytes, filename: str):
    from django.http import HttpResponse
    kind = "text/csv; charset=utf-8" if filename.endswith(".csv") else (
        "application/pdf" if filename.endswith(".pdf") else
        "application/zip" if filename.endswith(".zip") else XLSX_TYPE)
    res = HttpResponse(content, content_type=kind)
    res["Content-Disposition"] = f'attachment; filename="{filename}"'
    res["Access-Control-Expose-Headers"] = "Content-Disposition"
    res["Cache-Control"] = "no-store"
    return res


def export_file(fmt, title, stem, header, rows):
    """A ready file response for `rows` (lists of values in `header` order)."""
    if fmt not in EXPORT_FORMATS:
        from rest_framework.exceptions import ValidationError
        raise ValidationError({"export": "Use xlsx or csv."})
    content = to_xlsx(title, header, rows) if fmt == "xlsx" else to_csv(header, rows)
    return file_response(content, f"{stem}.{fmt}")


# ----------------------------------------------------------- row flattening

_ISO_DT = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}")
_ISO_D = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DECIMAL = re.compile(r"^-?(0|[1-9]\d*)\.\d+$")
# serializer fields that are never useful in a spreadsheet
_SKIP_KEYS = re.compile(r"(^|_)(url|urls|photo|photos|avatar|thumb|thumbnail|image|logo|signature|token|secret|password)s?$")
_NAME_KEYS = ("full_name", "name", "title", "label", "display_name", "username", "code", "email")


def _label(v):
    if isinstance(v, dict):
        for k in _NAME_KEYS:
            if v.get(k) not in (None, ""):
                return v[k]
        return ""
    return v


def cell(v):
    """One serialized value as a spreadsheet cell: names for nested objects, dates as dates, decimals as numbers."""
    if v is None:
        return None
    if isinstance(v, dict):
        return _label(v)
    if isinstance(v, (list, tuple)):
        return ", ".join(str(_label(x)) for x in v if _label(x) not in (None, ""))
    if isinstance(v, str):
        if _ISO_DT.match(v):
            from django.utils.dateparse import parse_datetime
            dt = parse_datetime(v)
            if dt is not None:
                return dt
        elif _ISO_D.match(v):
            try:
                return date.fromisoformat(v)
            except ValueError:
                return v
        elif _DECIMAL.match(v):
            return Decimal(v)
    return v


def dig(row, path):
    """`a.b.c` from a serialized row (dicts only)."""
    cur = row
    for part in path.split("."):
        if not isinstance(cur, dict):
            return None
        cur = cur.get(part)
    return cur


def auto_columns(sample: dict):
    """Columns for a serializer without its own list: every plain field, `x_detail` objects shown by name in
    place of their bare id, links/photos/secrets left out."""
    cols = []
    keys = list(sample.keys())
    for k in keys:
        if _SKIP_KEYS.search(k) or k in ("breakdown", "meta", "extra", "data", "payload"):
            continue
        if f"{k}_detail" in sample or f"{k}_name" in sample:
            continue   # the id; its detail/name column follows
        header = re.sub(r"_(detail|name)$", "", k) if k.endswith(("_detail", "_name")) and k != "name" else k
        header = header.replace("_", " ").strip().capitalize() or k
        if header.lower() == "id":
            header = "ID"
        cols.append((header, k))
    return cols


def rows_for(data, columns):
    """`data`: serialized rows. `columns`: [(header, "dotted.key" | callable(row))]."""
    header = [h for h, _ in columns]
    out = []
    for r in data:
        out.append([cell(get(r) if callable(get) else dig(r, get)) for _h, get in columns])
    return header, out


def export_format(request):
    fmt = (request.query_params.get("export") if hasattr(request, "query_params") else request.GET.get("export")) or ""
    return fmt if fmt in EXPORT_FORMATS else None


class ExportMixin:
    """Adds `?export=xlsx|csv` to a DRF list view: all matching rows (≤ EXPORT_MAX) instead of one page.

    Filters, search, sort and permissions are the list's own (`get_queryset` / `filter_queryset`), so the file has
    exactly what the screen shows, page after page. Set `export_columns` = [(header, key-or-callable)] for nice
    columns; otherwise every plain serializer field is written. `export_name` names the file and sheet.
    """
    export_columns = None
    export_name = None

    def _export_fmt(self):
        req = getattr(self, "request", None)
        return export_format(req) if req is not None and req.method == "GET" and getattr(self, "action", None) == "list" else None

    def paginate_queryset(self, queryset):
        if self._export_fmt():
            return list(queryset[:EXPORT_MAX]) if hasattr(queryset, "__getitem__") else list(queryset)[:EXPORT_MAX]
        return super().paginate_queryset(queryset)

    def get_paginated_response(self, data):
        if self._export_fmt():
            from rest_framework.response import Response
            return Response(data)
        return super().get_paginated_response(data)

    def get_export_columns(self, sample):
        return self.export_columns or auto_columns(sample or {})

    def finalize_response(self, request, response, *args, **kwargs):
        fmt = self._export_fmt()
        if fmt and getattr(response, "status_code", 0) == 200 and hasattr(response, "data"):
            data = response.data
            if isinstance(data, dict) and isinstance(data.get("results"), list):
                data = data["results"]
            if isinstance(data, list):
                data = data[:EXPORT_MAX]
                header, rows = rows_for(data, self.get_export_columns(data[0] if data else {}))
                name = self.export_name or getattr(self, "basename", None) or "export"
                response = export_file(fmt, name.replace("-", " ").capitalize(), name, header, rows)
        return super().finalize_response(request, response, *args, **kwargs)
